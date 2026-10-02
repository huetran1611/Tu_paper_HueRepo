"""Additional Saturday and Sunday departure experiments for the selected datasets."""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
from pathlib import Path

import hanoi_23_departure_experiment as departure
import hanoi_case_study_33_cloud as profiles
from hanoi_selected_instances import DATASETS, INSTANCE_INDICES

WEEKEND_PROFILES = ('thu7', 'chunhat')


def all_tasks():
    return [
        {'instance_index': index, 'actual_profile': profile, 'start_hour': hour, 'seed': seed}
        for profile in WEEKEND_PROFILES
        for index in INSTANCE_INDICES
        for hour in departure.START_HOURS
        for seed in departure.SEEDS
    ]


def build_matrix():
    tasks = all_tasks()
    base, extra = divmod(len(tasks), departure.JOB_COUNT)
    cursor = 0
    include = []
    for index in range(departure.JOB_COUNT):
        count = base + (index < extra)
        batch = tasks[cursor:cursor + count]
        cursor += count
        include.append({'batch_id': f'batch-{index + 1:03d}', 'task_count': count,
                        'tasks_json': json.dumps([[t['instance_index'], t['actual_profile'], t['start_hour'], t['seed']] for t in batch], separators=(',', ':'))})
    assert cursor == len(tasks)
    return {'include': include}


def run_batch(args):
    tasks = json.loads(args.tasks_json)
    allowed = {(t['instance_index'], t['actual_profile'], t['start_hour'], t['seed']) for t in all_tasks()}
    failures = []
    for index, profile, hour, seed in tasks:
        if (index, profile, hour, seed) not in allowed:
            raise ValueError(f'Unsupported task: {(index, profile, hour, seed)}')
        dataset = departure.instance_name(index)
        profiles.build_dataset_profiles(args.repo, args.profile_root, args.work_root, dataset)
        try:
            departure.run_job(args.repo, args.binary, index, hour,
                              args.output / profile / args.batch_id,
                              args.time_limit, args.segment_iters, seeds=(seed,),
                              speed_source=profiles.speed_path(args.work_root, dataset, profile),
                              actual_profile=profile)
        except (SystemExit, subprocess.TimeoutExpired) as error:
            failures.append(f'{profile}/{dataset}/{hour}/{seed}: {error}')
    if failures:
        raise SystemExit('\n'.join(failures))


def aggregate(args):
    from openpyxl import Workbook, load_workbook

    expected = {(t['instance_index'], t['actual_profile'], t['start_hour'], t['seed']) for t in all_tasks()}
    found = set()
    for path in args.input_root.rglob('result.json'):
        row = json.loads(path.read_text())
        key = (int(row['instance_index']), row['actual_profile'], int(row['start_hour']), int(row['seed']))
        if key not in expected or key in found:
            raise ValueError(f'Unexpected or duplicate result: {key}')
        found.add(key)
    failures = []
    for profile in WEEKEND_PROFILES:
        # Artifact directories have a profile layer between the batch and runs.
        staging = args.output / '_collected' / profile
        for path in args.input_root.rglob('result.json'):
            row = json.loads(path.read_text())
            if row['actual_profile'] == profile:
                target = staging / row['task_id'] / 'result.json'
                target.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(path, target)
        try:
            departure.aggregate(staging, args.output / profile)
        except SystemExit as error:
            failures.append(f'{profile}: {error}')
    summary = Workbook()
    summary.remove(summary.active)
    for title in ('By_Hour', 'By_Instance_Hour', 'All_Runs', 'Final_Solutions'):
        combined = []
        for profile in WEEKEND_PROFILES:
            path = args.output / profile / 'hanoi_20_departure_10seeds_results.xlsx'
            book = load_workbook(path, read_only=True, data_only=True)
            rows = iter(book[title].values)
            headers = next(rows)
            for row in rows:
                record = dict(zip(headers, row))
                combined.append({'actual_profile': profile, **record})
            book.close()
        departure.add_sheet(summary, title, combined)
    departure.add_sheet(summary, 'Validation', [{'check': 'complete_task_grid', 'found': len(found), 'expected': len(expected), 'passed': found == expected}])
    departure.add_sheet(summary, 'Config', [{'parameter': 'instances', 'value': ', '.join(DATASETS)}, {'parameter': 'profiles', 'value': ', '.join(WEEKEND_PROFILES)}, {'parameter': 'start_hours', 'value': '7,8,9,10'}, {'parameter': 'seeds', 'value': '1..10'}])
    summary.save(args.output / 'hanoi_20_weekend_departure_10seeds_results.xlsx')
    shutil.rmtree(args.output / '_collected')
    if failures or found != expected:
        raise SystemExit('\n'.join(failures + [f'Collected {len(found)}/{len(expected)} tasks']))


def main():
    parser = argparse.ArgumentParser()
    commands = parser.add_subparsers(dest='command', required=True)
    commands.add_parser('matrix')
    run = commands.add_parser('run-batch')
    run.add_argument('--repo', type=Path, default=Path.cwd())
    run.add_argument('--binary', type=Path, required=True)
    run.add_argument('--profile-root', type=Path, required=True)
    run.add_argument('--work-root', type=Path, required=True)
    run.add_argument('--tasks-json', required=True)
    run.add_argument('--batch-id', required=True)
    run.add_argument('--output', type=Path, required=True)
    run.add_argument('--time-limit', type=int, default=1800)
    run.add_argument('--segment-iters', type=int, default=100)
    agg = commands.add_parser('aggregate')
    agg.add_argument('--input-root', type=Path, required=True)
    agg.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if args.command == 'matrix':
        print(json.dumps(build_matrix(), separators=(',', ':')))
    elif args.command == 'run-batch':
        run_batch(args)
    else:
        aggregate(args)


if __name__ == '__main__':
    main()
