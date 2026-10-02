#!/usr/bin/env python3
"""GitHub Actions runner for the selected 20-instance Hanoi experiment."""

from __future__ import annotations

import argparse
import csv
import json
import re
import statistics
import subprocess
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from hanoi_selected_instances import DATA_ROOT, INSTANCE_INDICES


INSTANCE_COUNT = len(INSTANCE_INDICES)
JOB_COUNT = 250
START_HOURS = (7, 8, 9, 10)
SEEDS = tuple(range(1, 11))
MAX_ITERATIONS = 9_000
FIRST_TRAFFIC_HOUR = 6
LAST_TRAFFIC_HOUR = 17
EXPECTED_SPEED_ENTRIES = 101 * 101 * 12


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def instance_name(index: int) -> str:
    return f"set_{index:02d}"


def source_for(repo: Path, index: int) -> tuple[Path, str, str]:
    if index not in INSTANCE_INDICES:
        raise ValueError(f"Instance {index} is not in the selected dataset")
    dataset = instance_name(index)
    return repo / DATA_ROOT / dataset, f"hanoi_10x10_100_{dataset}_weekday", dataset


def companion_paths(repo: Path, index: int) -> dict[str, Path | str]:
    directory, stem, source_set = source_for(repo, index)
    base = directory / stem
    return {
        "directory": directory,
        "stem": stem,
        "source_set": source_set,
        "instance": Path(f"{base}.txt"),
        "truck": Path(f"{base}.truck_distance_m.txt"),
        "drone": Path(f"{base}.drone_distance_m.txt"),
        "speed": Path(f"{base}.v_ijl_kph.txt"),
    }


def all_tasks() -> list[dict[str, object]]:
    tasks = []
    for index in INSTANCE_INDICES:
        for hour in START_HOURS:
            for seed in SEEDS:
                tasks.append(
                    {
                        "task_id": (
                            f"{instance_name(index)}-{hour:02d}h-seed{seed:02d}"
                        ),
                        "instance_index": index,
                        "start_hour": hour,
                        "seed": seed,
                    }
                )
    return tasks


def build_matrix() -> dict[str, list[dict[str, object]]]:
    tasks = all_tasks()
    base_size, larger_jobs = divmod(len(tasks), JOB_COUNT)
    include = []
    cursor = 0
    for job_index in range(JOB_COUNT):
        task_count = base_size + (1 if job_index < larger_jobs else 0)
        job_tasks = tasks[cursor : cursor + task_count]
        cursor += task_count
        include.append(
            {
                "batch_id": f"batch-{job_index + 1:03d}",
                "task_count": task_count,
                "tasks_json": json.dumps(job_tasks, separators=(",", ":")),
            }
        )
    if cursor != len(tasks):
        raise AssertionError(f"Assigned {cursor}/{len(tasks)} tasks")
    return {"include": include}


def speed_hours(path: Path) -> tuple[set[int], int]:
    hours: set[int] = set()
    count = 0
    with path.open(encoding="utf-8") as stream:
        for line in stream:
            if not line.strip() or line.lstrip().startswith("#"):
                continue
            fields = line.split()
            hours.add(int(fields[2]))
            count += 1
    return hours, count


def validate_data(repo: Path) -> None:
    problems = []
    for index in INSTANCE_INDICES:
        paths = companion_paths(repo, index)
        for key in ("instance", "truck", "drone", "speed"):
            path = Path(paths[key])
            if not path.is_file() or path.stat().st_size == 0:
                problems.append(f"{instance_name(index)} missing {key}: {path}")
        speed = Path(paths["speed"])
        if speed.is_file():
            hours, count = speed_hours(speed)
            if hours != set(range(6, 19)):
                problems.append(
                    f"{instance_name(index)} source speed hours={sorted(hours)}"
                )
            if count != 101 * 101 * 13:
                problems.append(
                    f"{instance_name(index)} source speed entries={count}"
                )
    if problems:
        raise SystemExit("\n".join(problems))
    print(f"Validated {INSTANCE_COUNT} source instances with hours 6..18")


def write_trimmed_speed(source: Path, destination: Path) -> None:
    count = 0
    observed_hours: set[int] = set()
    destination.parent.mkdir(parents=True, exist_ok=True)
    with source.open(encoding="utf-8") as src, destination.open(
        "w", encoding="utf-8"
    ) as dst:
        dst.write("# i j hour v_ijl_kph; 12 intervals from 6h to 18h\n")
        dst.write("# Hours 6..17 are used; 18h is the exclusive end boundary\n")
        for line in src:
            if not line.strip() or line.lstrip().startswith("#"):
                continue
            i, j, hour_text, speed = line.split()[:4]
            hour = int(hour_text)
            if FIRST_TRAFFIC_HOUR <= hour <= LAST_TRAFFIC_HOUR:
                dst.write(f"{i} {j} {hour} {speed}\n")
                observed_hours.add(hour)
                count += 1
    if observed_hours != set(range(6, 18)) or count != EXPECTED_SPEED_ENTRIES:
        raise ValueError(
            f"Trimmed speed validation failed: hours={sorted(observed_hours)}, "
            f"entries={count}"
        )


def capture(pattern: str, text: str) -> str:
    matches = re.findall(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
    return matches[-1] if matches else ""


def parse_solution(path: Path) -> dict[str, object]:
    text = path.read_text(encoding="utf-8", errors="replace") if path.is_file() else ""
    initial = capture(r"Initial solution cost:\s*([-+0-9.eE]+)", text)
    improved = capture(r"Improved solution cost:\s*([-+0-9.eE]+)", text)
    worst = capture(r"Worst solution cost[^:]*:\s*([-+0-9.eE]+)", text)
    mean = capture(r"Mean solution cost[^:]*:\s*([-+0-9.eE]+)", text)
    elapsed = capture(r"Mean elapsed time:\s*([-+0-9.eE]+)", text)
    feasibility = capture(
        r"Final solution feasibility:\s*(FEASIBLE|INFEASIBLE)", text
    )
    return {
        "initial_makespan_s": float(initial) if initial else None,
        "improved_makespan_s": float(improved) if improved else None,
        "worst_makespan_s": float(worst) if worst else None,
        "mean_makespan_s": float(mean) if mean else None,
        "solver_elapsed_s": float(elapsed) if elapsed else None,
        "feasibility": feasibility or "UNKNOWN",
        "final_solution_text": text,
    }


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers and key != "final_solution_text":
                headers.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def run_job(
    repo: Path,
    binary: Path,
    index: int,
    start_hour: int,
    output: Path,
    time_limit: int,
    segment_iterations: int,
    seeds: tuple[int, ...] = SEEDS,
    speed_source: Path | None = None,
    actual_profile: str = "weekday",
) -> None:
    if start_hour not in START_HOURS:
        raise ValueError(f"Unsupported start hour: {start_hour}")
    paths = companion_paths(repo, index)
    dataset = instance_name(index)
    prefix = "" if actual_profile == "weekday" else f"{actual_profile}-"
    batch_id = f"{prefix}{dataset}-{start_hour:02d}h"
    batch_dir = output / batch_id
    speed = batch_dir / "input" / f"{actual_profile}_6h_17h.v_ijl_kph.txt"
    write_trimmed_speed(speed_source or Path(paths["speed"]), speed)
    results: list[dict[str, object]] = []

    for position, seed in enumerate(seeds, start=1):
        task_id = f"{prefix}{dataset}-{start_hour:02d}h-seed{seed:02d}"
        run_dir = batch_dir / "runs" / f"seed_{seed:02d}"
        run_dir.mkdir(parents=True, exist_ok=True)
        command = [
            str(binary.resolve()),
            str(Path(paths["instance"]).resolve()),
            f"--truck-distance-file={Path(paths['truck']).resolve()}",
            f"--drone-distance-file={Path(paths['drone']).resolve()}",
            f"--truck-vijl-file={speed.resolve()}",
            f"--start-hour={start_hour}",
            "--attempts=1",
            f"--iters={MAX_ITERATIONS}",
            f"--segment-iters={segment_iterations}",
            f"--time-limit={time_limit}",
            f"--seed={seed}",
            "--auto-tune",
        ]
        started_at = utc_now()
        started = time.monotonic()
        with (run_dir / "solver.log").open("w", encoding="utf-8") as log:
            completed = subprocess.run(
                command,
                cwd=run_dir,
                stdout=log,
                stderr=subprocess.STDOUT,
                check=False,
                timeout=time_limit + 600,
            )
        wall_time = time.monotonic() - started
        solution = run_dir / "output_solution_best.txt"
        parsed = parse_solution(solution)
        log_text = (run_dir / "solver.log").read_text(
            encoding="utf-8", errors="replace"
        )
        segments = log_text.count("=== End of Segment")
        if segments >= (MAX_ITERATIONS + segment_iterations - 1) // segment_iterations:
            termination = "ITERATION_LIMIT"
        elif wall_time >= time_limit - 2:
            termination = "TIME_LIMIT"
        else:
            termination = "EARLY_OR_ERROR"
        row: dict[str, object] = {
            "task_id": task_id,
            "instance": dataset,
            "actual_profile": actual_profile,
            "instance_index": index,
            "source_set": paths["source_set"],
            "start_hour": start_hour,
            "seed": seed,
            "status": "SUCCESS" if completed.returncode == 0 else "FAILED",
            "return_code": completed.returncode,
            "feasibility": parsed["feasibility"],
            "initial_makespan_s": parsed["initial_makespan_s"],
            "improved_makespan_s": parsed["improved_makespan_s"],
            "improved_makespan_h": (
                float(parsed["improved_makespan_s"]) / 3600.0
                if parsed["improved_makespan_s"] is not None
                else None
            ),
            "worst_makespan_s": parsed["worst_makespan_s"],
            "mean_makespan_s": parsed["mean_makespan_s"],
            "solver_elapsed_s": parsed["solver_elapsed_s"],
            "wall_time_s": wall_time,
            "termination_reason": termination,
            "executed_segments": segments,
            "max_iterations": MAX_ITERATIONS,
            "segment_iterations": segment_iterations,
            "time_limit_s": time_limit,
            "traffic_first_hour": FIRST_TRAFFIC_HOUR,
            "traffic_last_used_hour": LAST_TRAFFIC_HOUR,
            "traffic_segments": 12,
            "started_at": started_at,
            "finished_at": utc_now(),
            "solution_file": str(solution.relative_to(batch_dir)),
            "final_solution_text": parsed["final_solution_text"],
        }
        (run_dir / "result.json").write_text(
            json.dumps(row, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        results.append(row)
        write_csv(batch_dir / "summary.csv", results)
        print(
            f"[{position}/{len(seeds)}] {task_id}: {row['status']} "
            f"{row['feasibility']} ({wall_time:.2f}s)",
            flush=True,
        )

    metadata = {
        "batch_id": batch_id,
        "instance": dataset,
        "actual_profile": actual_profile,
        "start_hour": start_hour,
        "seeds": list(seeds),
        "traffic_hours_used": list(range(6, 18)),
        "traffic_18h_used": False,
        "max_iterations": MAX_ITERATIONS,
        "segment_iterations": segment_iterations,
        "time_limit_s": time_limit,
    }
    (batch_dir / "metadata.json").write_text(
        json.dumps(metadata, indent=2) + "\n", encoding="utf-8"
    )
    speed.unlink(missing_ok=True)
    speed.parent.rmdir()
    failed = [row for row in results if row["status"] != "SUCCESS"]
    if failed:
        raise SystemExit(f"Failed seeds in {batch_id}: {len(failed)}")


def run_batch(
    repo: Path,
    binary: Path,
    batch_id: str,
    tasks: list[dict[str, object]],
    output: Path,
    time_limit: int,
    segment_iterations: int,
) -> None:
    batch_dir = output / batch_id
    batch_dir.mkdir(parents=True, exist_ok=True)
    failures = []
    for position, task in enumerate(tasks, start=1):
        print(
            f"Batch {batch_id}: task {position}/{len(tasks)} "
            f"({task['task_id']})",
            flush=True,
        )
        try:
            run_job(
                repo,
                binary,
                int(task["instance_index"]),
                int(task["start_hour"]),
                batch_dir,
                time_limit,
                segment_iterations,
                seeds=(int(task["seed"]),),
            )
        except SystemExit as error:
            failures.append({"task_id": task["task_id"], "error": str(error)})
    (batch_dir / "batch_metadata.json").write_text(
        json.dumps(
            {"batch_id": batch_id, "task_count": len(tasks), "tasks": tasks},
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    if failures:
        raise SystemExit(
            f"{len(failures)} task(s) failed in {batch_id}: "
            + ", ".join(str(item["task_id"]) for item in failures)
        )


def add_sheet(workbook, name: str, rows: list[dict[str, object]]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(name)
    if not rows:
        sheet.append(["No data"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, header in enumerate(headers, start=1):
        width = 80 if header == "final_solution_text" else min(max(len(header) + 3, 13), 32)
        sheet.column_dimensions[get_column_letter(column)].width = width


def aggregate(input_root: Path, output: Path) -> None:
    from openpyxl import Workbook

    result_files = sorted(input_root.rglob("result.json"))
    rows = [json.loads(path.read_text(encoding="utf-8")) for path in result_files]
    unique = {str(row["task_id"]): row for row in rows}
    rows = sorted(
        unique.values(),
        key=lambda row: (
            int(row["instance_index"]),
            int(row["start_hour"]),
            int(row["seed"]),
        ),
    )

    configuration_rows = []
    config_groups: dict[tuple[str, int], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        config_groups[(str(row["instance"]), int(row["start_hour"]))].append(row)
    for (instance, hour), group in sorted(config_groups.items()):
        values = [
            float(row["improved_makespan_s"])
            for row in group
            if row.get("improved_makespan_s") is not None
        ]
        configuration_rows.append(
            {
                "instance": instance,
                "start_hour": hour,
                "runs": len(group),
                "feasible_runs": sum(row.get("feasibility") == "FEASIBLE" for row in group),
                "mean_makespan_s": statistics.fmean(values) if values else None,
                "mean_makespan_h": statistics.fmean(values) / 3600.0 if values else None,
                "best_makespan_s": min(values) if values else None,
                "worst_makespan_s": max(values) if values else None,
                "stdev_makespan_s": statistics.stdev(values) if len(values) > 1 else 0.0,
                "mean_wall_time_s": statistics.fmean(float(row["wall_time_s"]) for row in group),
            }
        )

    hour_rows = []
    for hour in START_HOURS:
        configurations = [
            row for row in configuration_rows if int(row["start_hour"]) == hour
        ]
        means = [float(row["mean_makespan_s"]) for row in configurations]
        hour_rows.append(
            {
                "start_hour": hour,
                "instances": len(configurations),
                "runs": sum(int(row["runs"]) for row in configurations),
                "mean_makespan_s": statistics.fmean(means) if means else None,
                "mean_makespan_h": statistics.fmean(means) / 3600.0 if means else None,
                "stdev_across_instance_means_s": statistics.stdev(means) if len(means) > 1 else 0.0,
            }
        )
    valid_hour_rows = [row for row in hour_rows if row["mean_makespan_s"] is not None]
    best_hour = min(
        valid_hour_rows, key=lambda row: float(row["mean_makespan_s"])
    )["start_hour"] if valid_hour_rows else None

    solution_rows = []
    for row in rows:
        solution_rows.append(
            {
                "task_id": row["task_id"],
                "instance": row["instance"],
                "start_hour": row["start_hour"],
                "seed": row["seed"],
                "feasibility": row["feasibility"],
                "improved_makespan_s": row["improved_makespan_s"],
                "final_solution_text": row.get("final_solution_text", ""),
            }
        )
    run_rows = [
        {key: value for key, value in row.items() if key != "final_solution_text"}
        for row in rows
    ]
    validation = [
        {"check": "expected_runs", "value": INSTANCE_COUNT * len(START_HOURS) * len(SEEDS)},
        {"check": "collected_runs", "value": len(rows)},
        {"check": "successful_runs", "value": sum(row.get("status") == "SUCCESS" for row in rows)},
        {"check": "feasible_runs", "value": sum(row.get("feasibility") == "FEASIBLE" for row in rows)},
        {"check": "configuration_groups", "value": len(configuration_rows)},
        {"check": "best_start_hour", "value": best_hour},
        {"check": "traffic_hours_used", "value": "6,7,8,9,10,11,12,13,14,15,16,17"},
        {"check": "traffic_18h_used", "value": False},
    ]
    config = [
        {"parameter": "instances", "value": INSTANCE_COUNT},
        {"parameter": "start_hours", "value": "7,8,9,10"},
        {"parameter": "seeds", "value": "1,2,3,4,5,6,7,8,9,10"},
        {"parameter": "max_iterations", "value": MAX_ITERATIONS},
        {"parameter": "traffic_segments", "value": 12},
        {"parameter": "traffic_end_boundary", "value": "18h exclusive"},
    ]

    output.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "By_Hour", hour_rows)
    add_sheet(workbook, "By_Instance_Hour", configuration_rows)
    add_sheet(workbook, "All_Runs", run_rows)
    add_sheet(workbook, "Final_Solutions", solution_rows)
    add_sheet(workbook, "Config", config)
    add_sheet(workbook, "Validation", validation)
    workbook_path = output / "hanoi_20_departure_10seeds_results.xlsx"
    workbook.save(workbook_path)
    write_csv(output / "all_runs.csv", run_rows)
    write_csv(output / "by_instance_hour.csv", configuration_rows)
    write_csv(output / "by_hour.csv", hour_rows)
    print(f"Workbook: {workbook_path}")
    expected_runs = INSTANCE_COUNT * len(START_HOURS) * len(SEEDS)
    print(f"Collected: {len(rows)}/{expected_runs}; best hour: {best_hour}")
    if len(rows) != expected_runs:
        raise SystemExit(
            f"Expected {expected_runs} unique runs, found {len(rows)}"
        )


def main() -> None:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)

    subparsers.add_parser("matrix")
    validate_parser = subparsers.add_parser("validate-data")
    validate_parser.add_argument("--repo", type=Path, default=Path.cwd())

    run_parser = subparsers.add_parser("run-job")
    run_parser.add_argument("--repo", type=Path, default=Path.cwd())
    run_parser.add_argument("--binary", type=Path, required=True)
    run_parser.add_argument("--instance-index", type=int, required=True)
    run_parser.add_argument("--start-hour", type=int, required=True)
    run_parser.add_argument("--output", type=Path, required=True)
    run_parser.add_argument("--time-limit", type=int, default=1_800)
    run_parser.add_argument("--segment-iters", type=int, default=100)

    batch_parser = subparsers.add_parser("run-batch")
    batch_parser.add_argument("--repo", type=Path, default=Path.cwd())
    batch_parser.add_argument("--binary", type=Path, required=True)
    batch_parser.add_argument("--batch-id", required=True)
    batch_parser.add_argument("--tasks-json", required=True)
    batch_parser.add_argument("--output", type=Path, required=True)
    batch_parser.add_argument("--time-limit", type=int, default=1_800)
    batch_parser.add_argument("--segment-iters", type=int, default=100)

    aggregate_parser = subparsers.add_parser("aggregate")
    aggregate_parser.add_argument("--input-root", type=Path, required=True)
    aggregate_parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    if args.command == "matrix":
        print(json.dumps(build_matrix(), separators=(",", ":")))
    elif args.command == "validate-data":
        validate_data(args.repo.resolve())
    elif args.command == "run-job":
        run_job(
            args.repo.resolve(),
            args.binary,
            args.instance_index,
            args.start_hour,
            args.output,
            args.time_limit,
            args.segment_iters,
        )
    elif args.command == "run-batch":
        tasks = json.loads(args.tasks_json)
        if not isinstance(tasks, list) or not tasks:
            raise ValueError("--tasks-json must contain a non-empty JSON list")
        run_batch(
            args.repo.resolve(),
            args.binary,
            args.batch_id,
            tasks,
            args.output,
            args.time_limit,
            args.segment_iters,
        )
    elif args.command == "aggregate":
        aggregate(args.input_root, args.output)


if __name__ == "__main__":
    main()
