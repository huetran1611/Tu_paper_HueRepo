#!/usr/bin/env python3
"""Run the selected 20-instance Hanoi P0/P1/P2 experiment on GitHub Actions."""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import statistics
import subprocess
import time
import unicodedata
from collections import defaultdict
from pathlib import Path

from hanoi_selected_instances import DATA_ROOT, DATASETS


INSTANCE_COUNT = len(DATASETS)
PROFILES = ("weekday", "thu7", "chunhat")
START_HOURS = (7, 8, 9, 10)
TRAFFIC_HOURS = tuple(range(6, 18))
SEEDS = tuple(range(1, 11))
MAX_ITERATIONS = 9_000
SEGMENT_ITERATIONS = 300
TIME_LIMIT_SECONDS = 600
REPLAY_TIMEOUT_SECONDS = 120
JOB_COUNT = 250

PROFILE_FILENAMES = {
    "weekday": "hanoi_traffic_speed_thu2_weekday_representative_10pct.csv",
    "thu7": "hanoi_traffic_speed_thu7_complete_10pct.csv",
    "chunhat": "hanoi_traffic_speed_chunhat_complete.csv",
}


def dataset_index(dataset: str) -> int:
    return int(dataset.split("_")[1])


def source_paths(repo: Path, dataset: str) -> dict[str, Path]:
    if dataset not in DATASETS:
        raise ValueError(f"Dataset {dataset} is not in the selected dataset")
    directory = repo / DATA_ROOT / dataset
    stem = f"hanoi_10x10_100_{dataset}_weekday"
    return {
        "instance": directory / f"{stem}.txt",
        "truck": directory / f"{stem}.truck_distance_m.txt",
        "drone": directory / f"{stem}.drone_distance_m.txt",
        "points": directory / "nga_tu_so_10x10km_100_points_min500m.csv",
    }


def normalize(value: str) -> str:
    value = value.casefold().replace("đ", "d")
    decomposed = unicodedata.normalize("NFD", value)
    return " ".join(
        "".join(
            character
            for character in decomposed
            if unicodedata.category(character) != "Mn"
        ).split()
    )


def load_road_profiles(path: Path) -> dict[tuple[str, str], tuple[float, ...]]:
    profiles = {}
    with path.open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            key = (normalize(row["Ten_Duong"]), normalize(row["Loai_Duong"]))
            profiles[key] = tuple(float(row[f"{hour}h"]) for hour in TRAFFIC_HOURS)
    return profiles


def load_node_profiles(
    points_path: Path, road_profiles: dict[tuple[str, str], tuple[float, ...]]
) -> list[tuple[tuple[str, str] | None, tuple[float, ...]]]:
    with points_path.open(encoding="utf-8-sig", newline="") as stream:
        points = sorted(csv.DictReader(stream), key=lambda row: int(row["point_id"]))
    if len(points) != 100:
        raise ValueError(f"Expected 100 points in {points_path}, found {len(points)}")
    customer_nodes = []
    for point in points:
        key = (normalize(point["road_name"]), normalize(point["district"]))
        if key not in road_profiles:
            raise KeyError(f"Missing road profile {key} for {points_path}")
        customer_nodes.append((key, road_profiles[key]))
    depot = tuple(
        statistics.fmean(node[1][segment] for node in customer_nodes)
        for segment in range(len(TRAFFIC_HOURS))
    )
    return [(None, depot), *customer_nodes]


def edge_speed(
    source: tuple[tuple[str, str] | None, tuple[float, ...]],
    destination: tuple[tuple[str, str] | None, tuple[float, ...]],
    segment: int,
) -> float:
    source_key, source_profile = source
    destination_key, destination_profile = destination
    if source_key is not None and source_key == destination_key:
        return source_profile[segment]
    return (source_profile[segment] + destination_profile[segment]) / 2.0


def speed_path(work_root: Path, dataset: str, label: str) -> Path:
    return work_root / dataset / f"{label}.v_ijl_kph.txt"


def build_dataset_profiles(
    repo: Path, profile_root: Path, work_root: Path, dataset: str,
    aggregation: str = "arithmetic",
) -> None:
    marker = work_root / dataset / "complete.json"
    if marker.is_file():
        metadata = json.loads(marker.read_text(encoding="utf-8"))
        if metadata.get("aggregation", "arithmetic") == aggregation:
            return
    paths = source_paths(repo, dataset)
    for path in paths.values():
        if not path.is_file():
            raise FileNotFoundError(path)
    node_profiles = {
        profile: load_node_profiles(
            paths["points"],
            load_road_profiles(profile_root / PROFILE_FILENAMES[profile]),
        )
        for profile in PROFILES
    }
    output_dir = work_root / dataset
    output_dir.mkdir(parents=True, exist_ok=True)
    handles = {
        profile: speed_path(work_root, dataset, profile).open("w", encoding="utf-8")
        for profile in PROFILES
    }
    p1_handle = speed_path(work_root, dataset, "P1").open("w", encoding="utf-8")
    p0_handle = (
        speed_path(work_root, dataset, "P0").open("w", encoding="utf-8")
        if aggregation == "harmonic" else None
    )
    total_speed = 0.0
    total_count = 0
    try:
        for handle in [*handles.values(), p1_handle]:
            handle.write("# i j hour v_ijl_kph; hours 6..17, 18h exclusive\n")
        if p0_handle is not None:
            p0_handle.write("# arc-wise harmonic speed across days/hours; 18h exclusive\n")
        for i in range(101):
            for j in range(101):
                edge_values = []
                for segment, hour in enumerate(TRAFFIC_HOURS):
                    values = {
                        profile: edge_speed(
                            node_profiles[profile][i], node_profiles[profile][j], segment
                        )
                        for profile in PROFILES
                    }
                    if any(value <= 0 for value in values.values()):
                        raise ValueError(f"Nonpositive speed for {dataset}, {i}, {j}, {hour}")
                    edge_values.extend(values.values())
                    for profile, value in values.items():
                        handles[profile].write(f"{i} {j} {hour} {value:.6f}\n")
                    p1_speed = (
                        statistics.harmonic_mean(values.values())
                        if aggregation == "harmonic" else statistics.fmean(values.values())
                    )
                    p1_handle.write(f"{i} {j} {hour} {p1_speed:.6f}\n")
                    total_speed += sum(values.values())
                    total_count += len(values)
                if p0_handle is not None:
                    # Pool days/hours for this arc, preserving spatial variation.
                    p0_speed = statistics.harmonic_mean(edge_values)
                    for hour in TRAFFIC_HOURS:
                        p0_handle.write(f"{i} {j} {hour} {p0_speed:.6f}\n")
    finally:
        for handle in handles.values():
            handle.close()
        p1_handle.close()
        if p0_handle is not None:
            p0_handle.close()

    p0_speed = total_speed / total_count
    if aggregation == "arithmetic":
        write_static_profile(work_root, dataset, p0_speed)
    marker.write_text(
        json.dumps(
            {
                "dataset": dataset,
                "traffic_hours": list(TRAFFIC_HOURS),
                "aggregation": aggregation,
                "speed_representation": "arc speeds derived from endpoint road profiles",
                "p0_static_speed_kph": p0_speed if aggregation == "arithmetic" else None,
                "p0": "arc-wise harmonic across days/hours" if aggregation == "harmonic" else "global arithmetic mean",
                "p1": f"edge/hour {aggregation} mean across weekday, Saturday, Sunday",
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )


def write_static_profile(work_root: Path, dataset: str, p0_speed: float) -> None:
    with speed_path(work_root, dataset, "P0").open("w", encoding="utf-8") as stream:
        stream.write(
            f"# i j hour v_ijl_kph; static mean={p0_speed:.6f}; "
            "hours 6..17, 18h exclusive\n"
        )
        for i in range(101):
            for j in range(101):
                for hour in TRAFFIC_HOURS:
                    stream.write(f"{i} {j} {hour} {p0_speed:.6f}\n")


def validate_data(repo: Path, profile_root: Path) -> None:
    missing = []
    for dataset in DATASETS:
        for label, path in source_paths(repo, dataset).items():
            if not path.is_file() or path.stat().st_size == 0:
                missing.append(f"{dataset} missing {label}: {path}")
    for filename in PROFILE_FILENAMES.values():
        path = profile_root / filename
        if not path.is_file() or path.stat().st_size == 0:
            missing.append(f"missing day profile: {path}")
    if missing:
        raise SystemExit("\n".join(missing))
    print(f"Validated {INSTANCE_COUNT} instances and {len(PROFILES)} day profiles")


def policy_tasks() -> dict[str, list[dict[str, object]]]:
    return {
        "P0": [
            {"task_id": f"P0-{dataset}-seed{seed:02d}", "policy": "P0", "dataset": dataset, "seed": seed}
            for dataset in DATASETS for seed in SEEDS
        ],
        "P1": [
            {"task_id": f"P1-{dataset}-{hour:02d}h-seed{seed:02d}", "policy": "P1", "dataset": dataset, "start_hour": hour, "seed": seed}
            for dataset in DATASETS for hour in START_HOURS for seed in SEEDS
        ],
        "P2": [
            {"task_id": f"P2-{dataset}-{profile}-{hour:02d}h-seed{seed:02d}", "policy": "P2", "dataset": dataset, "actual_profile": profile, "start_hour": hour, "seed": seed}
            for dataset in DATASETS for profile in PROFILES for hour in START_HOURS for seed in SEEDS
        ],
    }


def compact_task(task: dict[str, object]) -> list[object]:
    policy = str(task["policy"])
    common = [dataset_index(str(task["dataset"])), int(task["seed"])]
    if policy == "P0":
        return [0, *common]
    if policy == "P1":
        return [1, *common, int(task["start_hour"])]
    return [
        2,
        *common,
        PROFILES.index(str(task["actual_profile"])),
        int(task["start_hour"]),
    ]


def expand_task(values: list[object]) -> dict[str, object]:
    policy_id, index, seed = map(int, values[:3])
    dataset = f"set_{index:02d}"
    if policy_id == 0:
        return {
            "task_id": f"P0-{dataset}-seed{seed:02d}",
            "policy": "P0",
            "dataset": dataset,
            "seed": seed,
        }
    hour = int(values[-1])
    if policy_id == 1:
        return {
            "task_id": f"P1-{dataset}-{hour:02d}h-seed{seed:02d}",
            "policy": "P1",
            "dataset": dataset,
            "start_hour": hour,
            "seed": seed,
        }
    profile = PROFILES[int(values[3])]
    return {
        "task_id": f"P2-{dataset}-{profile}-{hour:02d}h-seed{seed:02d}",
        "policy": "P2",
        "dataset": dataset,
        "actual_profile": profile,
        "start_hour": hour,
        "seed": seed,
    }


def build_matrix(policy: str | None = None) -> dict[str, list[dict[str, object]]]:
    queues = policy_tasks()
    if policy is not None:
        tasks = queues[policy]
        # P2 has no cross-profile replays and fits in 240 batches of ten solves.
        batch_size = 10 if policy == "P2" else 4
        return {"include": [
            {
                "batch_id": f"{policy.lower()}-batch{index // batch_size + 1:03d}",
                "task_count": len(tasks[index:index + batch_size]),
                "replay_count": len(tasks[index:index + batch_size]) * {"P0": 12, "P1": 3, "P2": 0}[policy],
                "tasks_json": json.dumps([compact_task(task) for task in tasks[index:index + batch_size]], separators=(",", ":")),
            }
            for index in range(0, len(tasks), batch_size)
        ]}
    cursors = defaultdict(int)
    # (number of jobs, P0, P1, P2) gives 250 jobs and 3,400 runs.
    specifications = ((100, 1, 3, 10), (50, 1, 4, 9), (50, 1, 3, 9), (50, 0, 3, 10))
    include = []
    for count, p0_count, p1_count, p2_count in specifications:
        for _ in range(count):
            tasks = []
            for policy, take in (("P0", p0_count), ("P1", p1_count), ("P2", p2_count)):
                start = cursors[policy]
                tasks.extend(queues[policy][start : start + take])
                cursors[policy] += take
            index = len(include) + 1
            include.append(
                {
                    "batch_id": f"policy-batch{index:03d}",
                    "task_count": len(tasks),
                    "replay_count": p0_count * 12 + p1_count * 3,
                    "tasks_json": json.dumps(
                        [compact_task(task) for task in tasks], separators=(",", ":")
                    ),
                }
            )
    for policy, tasks in queues.items():
        if cursors[policy] != len(tasks):
            raise AssertionError(f"Assigned {cursors[policy]}/{len(tasks)} {policy} tasks")
    return {"include": include}


def build_single_trip_matrix() -> dict[str, list[dict[str, object]]]:
    tasks = [
        [dataset_index(dataset), PROFILES.index(profile), hour, seed]
        for dataset in DATASETS
        for profile in PROFILES
        for hour in START_HOURS
        for seed in SEEDS
    ]
    include = []
    cursor = 0
    base_size, larger_jobs = divmod(len(tasks), JOB_COUNT)
    for job_index in range(JOB_COUNT):
        task_count = base_size + (1 if job_index < larger_jobs else 0)
        batch = tasks[cursor : cursor + task_count]
        cursor += task_count
        include.append(
            {
                "batch_id": f"single-trip-batch{job_index + 1:03d}",
                "task_count": task_count,
                "tasks_json": json.dumps(batch, separators=(",", ":")),
            }
        )
    if cursor != len(tasks):
        raise AssertionError(f"Assigned {cursor}/{len(tasks)} single-trip tasks")
    return {"include": include}


def expand_single_trip_task(values: list[object]) -> dict[str, object]:
    index, profile_index, hour, seed = map(int, values)
    dataset = f"set_{index:02d}"
    profile = PROFILES[profile_index]
    return {
        "task_id": f"ST-{dataset}-{profile}-{hour:02d}h-seed{seed:02d}",
        "dataset": dataset,
        "profile": profile,
        "start_hour": hour,
        "seed": seed,
    }


def capture(pattern: str, text: str) -> str:
    matches = re.findall(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
    return matches[-1] if matches else ""


def parse_solution(path: Path) -> dict[str, object]:
    text = path.read_text(encoding="utf-8", errors="replace") if path.is_file() else ""
    makespan = capture(r"Improved solution cost:\s*([-+0-9.eE]+)", text)
    feasibility = capture(r"Final solution feasibility:\s*(FEASIBLE|INFEASIBLE)", text)
    validation = re.search(
        r"Total validation:\s*Makespan=([-+0-9.eE]+),\s*Deadline violation=([-+0-9.eE]+),\s*Energy violation=([-+0-9.eE]+),\s*Capacity violation=([-+0-9.eE]+)",
        text,
    )
    return {
        "makespan_s": float(makespan) if makespan else None,
        "feasibility": feasibility or "UNKNOWN",
        "deadline_violation": float(validation.group(2)) if validation else None,
        "energy_violation": float(validation.group(3)) if validation else None,
        "capacity_violation": float(validation.group(4)) if validation else None,
        "final_solution_text": text,
    }


def parse_route_metrics(solution_text: str) -> dict[str, object]:
    trip_counts = []
    customers_per_trip = []
    truck_customers = 0
    truck_time_s = 0.0
    for line in solution_text.splitlines():
        match = re.match(r"^Truck \d+:\s*(.*?)\|Truck Time:\s*([-+0-9.eE]+)", line)
        if not match:
            continue
        route = [int(value) for value in match.group(1).split()]
        trips = []
        current = []
        for node in route[1:]:
            if node == 0:
                if current:
                    trips.append(current)
                    current = []
            else:
                current.append(node)
        if current:
            trips.append(current)
        trip_counts.append(len(trips))
        customers_per_trip.append([len(trip) for trip in trips])
        truck_customers += sum(len(trip) for trip in trips)
        truck_time_s += float(match.group(2))
    drone_customers = 0
    for line in solution_text.splitlines():
        match = re.match(r"^Drone \d+:\s*(.*?)\|Drone Time:", line)
        if match:
            drone_customers += sum(int(value) != 0 for value in match.group(1).split())
    used_trip_counts = [count for count in trip_counts if count > 0]
    later_trip_customers = sum(
        sum(per_trip[1:]) for per_trip in customers_per_trip if len(per_trip) > 1
    )
    return {
        "truck_trip_counts": json.dumps(trip_counts, separators=(",", ":")),
        "truck_customers_per_trip": json.dumps(customers_per_trip, separators=(",", ":")),
        "truck_trips": sum(trip_counts),
        "used_trucks": len(used_trip_counts),
        "trips_per_used_truck": (
            statistics.fmean(used_trip_counts) if used_trip_counts else 0.0
        ),
        "intermediate_depot_returns": sum(max(0, count - 1) for count in trip_counts),
        "later_trip_customers": later_trip_customers,
        "later_trip_customers_pct": 100.0 * later_trip_customers / 100.0,
        "truck_customers": truck_customers,
        "drone_customers": drone_customers,
        "drone_customers_pct": 100.0 * drone_customers / 100.0,
        "total_truck_travel_and_service_time_s": truck_time_s,
    }


def run_command(
    command: list[str], cwd: Path, log_path: Path, timeout: int, allowed: tuple[int, ...]
) -> tuple[int, float, str]:
    started = time.monotonic()
    try:
        completed = subprocess.run(
            command,
            cwd=cwd,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            timeout=timeout,
            check=False,
        )
        output = completed.stdout
        return_code = completed.returncode
    except subprocess.TimeoutExpired as error:
        output = (error.stdout or "") + f"\nTimed out after {timeout} seconds\n"
        return_code = 124
    elapsed = time.monotonic() - started
    log_path.write_text("Command: " + " ".join(command) + "\n\n" + output, encoding="utf-8")
    if return_code not in allowed:
        raise RuntimeError(f"Command failed ({return_code}): {log_path}")
    return return_code, elapsed, output


def solver_command(
    binary: Path,
    source: dict[str, Path],
    speed: Path,
    start_hour: int,
    seed: int | None = None,
    *,
    iterations: int | None = None,
    segment_iterations: int | None = None,
    time_limit: int | None = None,
    auto_tune: bool = True,
) -> list[str]:
    command = [
        str(binary),
        str(source["instance"]),
        f"--truck-distance-file={source['truck']}",
        f"--drone-distance-file={source['drone']}",
        f"--truck-vijl-file={speed}",
        f"--start-hour={start_hour}",
    ]
    if seed is not None:
        command.extend(
            [
                "--attempts=1",
                f"--iters={iterations if iterations is not None else MAX_ITERATIONS}",
                f"--segment-iters={segment_iterations if segment_iterations is not None else SEGMENT_ITERATIONS}",
                f"--time-limit={time_limit if time_limit is not None else TIME_LIMIT_SECONDS}",
                f"--seed={seed}",
            ]
        )
        if auto_tune:
            command.append("--auto-tune")
    return command


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        return
    headers = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def evaluate_solution(
    binary: Path,
    source: dict[str, Path],
    speed: Path,
    solution: Path,
    policy: str,
    dataset: str,
    profile: str,
    hour: int,
    seed: int,
    output_dir: Path,
) -> dict[str, object]:
    output_dir.mkdir(parents=True, exist_ok=True)
    command = solver_command(binary, source, speed, hour)
    command.append(f"--evaluate-solution={solution}")
    return_code, elapsed, output = run_command(
        command,
        output_dir,
        output_dir / "evaluation.log",
        REPLAY_TIMEOUT_SECONDS,
        (0, 2),
    )
    parsed = parse_solution(output_dir / "output_solution_evaluated.txt")
    validation = re.search(
        r"Total validation:\s*Makespan=([-+0-9.eE]+),\s*Deadline violation=([-+0-9.eE]+),\s*Energy violation=([-+0-9.eE]+),\s*Capacity violation=([-+0-9.eE]+)",
        output,
    )
    return {
        "evaluation_id": f"{policy}-{dataset}-{profile}-{hour:02d}h-seed{seed:02d}",
        "policy": policy,
        "dataset": dataset,
        "actual_profile": profile,
        "start_hour": hour,
        "seed": seed,
        "return_code": return_code,
        "feasibility": parsed["feasibility"],
        "realized_makespan_s": parsed["makespan_s"],
        "deadline_violation": float(validation.group(2)) if validation else None,
        "energy_violation": float(validation.group(3)) if validation else None,
        "capacity_violation": float(validation.group(4)) if validation else None,
        "evaluation_wall_time_s": elapsed,
    }


def run_batch(
    repo: Path,
    binary: Path,
    profile_root: Path,
    work_root: Path,
    tasks: list[dict[str, object]],
    output: Path,
    aggregation: str = "arithmetic",
    iterations: int | None = None,
    segment_iterations: int | None = None,
    time_limit: int | None = None,
) -> None:
    iterations = MAX_ITERATIONS if iterations is None else iterations
    segment_iterations = SEGMENT_ITERATIONS if segment_iterations is None else segment_iterations
    time_limit = TIME_LIMIT_SECONDS if time_limit is None else time_limit
    if min(iterations, segment_iterations, time_limit) <= 0:
        raise ValueError("Search budgets must be positive")
    output.mkdir(parents=True, exist_ok=True)
    optimization_rows: list[dict[str, object]] = []
    evaluation_rows: list[dict[str, object]] = []
    failures = []
    for dataset in sorted({str(task["dataset"]) for task in tasks}):
        build_dataset_profiles(repo, profile_root, work_root, dataset, aggregation)
    (output / "experiment_config.json").write_text(
        json.dumps({"aggregation": aggregation, "iterations": iterations,
                    "segment_iterations": segment_iterations, "time_limit_seconds": time_limit,
                    "auto_tune": aggregation != "harmonic", "tasks": tasks}, indent=2) + "\n",
        encoding="utf-8",
    )

    for position, task in enumerate(tasks, start=1):
        policy = str(task["policy"])
        dataset = str(task["dataset"])
        seed = int(task["seed"])
        hour = 7 if policy == "P0" else int(task["start_hour"])
        task_id = str(task["task_id"])
        task_dir = output / task_id
        task_dir.mkdir(parents=True, exist_ok=True)
        source = source_paths(repo, dataset)
        speed_label = policy if policy in ("P0", "P1") else str(task["actual_profile"])
        speed = speed_path(work_root, dataset, speed_label)
        print(f"[{position}/{len(tasks)}] optimizing {task_id}", flush=True)
        try:
            command = solver_command(
                binary, source, speed, hour, seed,
                iterations=iterations, segment_iterations=segment_iterations,
                time_limit=time_limit, auto_tune=aggregation != "harmonic",
            )
            _, elapsed, _ = run_command(
                command,
                task_dir,
                task_dir / "solver.log",
                time_limit + 120,
                (0,),
            )
            solution = task_dir / "output_solution_best.txt"
            saved_solution = task_dir / f"final_solution_{task_id}.txt"
            shutil.copy2(solution, saved_solution)
            parsed = parse_solution(saved_solution)
            optimization_rows.append(
                {
                    "task_id": task_id,
                    "policy": policy,
                    "dataset": dataset,
                    "planned_profile": task.get("actual_profile", "expected"),
                    "planned_start_hour": hour,
                    "seed": seed,
                    "feasibility": parsed["feasibility"],
                    "planned_makespan_s": parsed["makespan_s"],
                    "wall_time_s": elapsed,
                    "solution_file": str(saved_solution.relative_to(output)),
                }
            )
            write_csv(output / "optimization_results.csv", optimization_rows)

            if policy == "P0":
                cases = ((profile, replay_hour) for profile in PROFILES for replay_hour in START_HOURS)
            elif policy == "P1":
                cases = ((profile, hour) for profile in PROFILES)
            else:
                profile = str(task["actual_profile"])
                evaluation_rows.append(
                    {
                        "evaluation_id": f"P2-{dataset}-{profile}-{hour:02d}h-seed{seed:02d}",
                        "policy": "P2",
                        "dataset": dataset,
                        "actual_profile": profile,
                        "start_hour": hour,
                        "seed": seed,
                        "return_code": 0,
                        "feasibility": parsed["feasibility"],
                        "realized_makespan_s": parsed["makespan_s"],
                        "deadline_violation": parsed["deadline_violation"],
                        "energy_violation": parsed["energy_violation"],
                        "capacity_violation": parsed["capacity_violation"],
                        "evaluation_wall_time_s": 0.0,
                    }
                )
                cases = ()
            for profile, replay_hour in cases:
                evaluated = evaluate_solution(
                    binary,
                    source,
                    speed_path(work_root, dataset, profile),
                    saved_solution,
                    policy,
                    dataset,
                    profile,
                    replay_hour,
                    seed,
                    task_dir / "replays" / profile / f"{replay_hour:02d}h",
                )
                if aggregation == "harmonic" and evaluated["feasibility"] != "FEASIBLE":
                    evaluated["realized_makespan_s"] = None
                evaluation_rows.append(evaluated)
                write_csv(output / "evaluation_results.csv", evaluation_rows)
            write_csv(output / "evaluation_results.csv", evaluation_rows)
        except Exception as error:
            failures.append({"task_id": task_id, "error": repr(error)})
            (task_dir / "failure.txt").write_text(repr(error) + "\n", encoding="utf-8")
            print(f"FAILED {task_id}: {error}", flush=True)
    (output / "failures.json").write_text(
        json.dumps(failures, indent=2) + "\n", encoding="utf-8"
    )
    if failures:
        raise SystemExit(f"{len(failures)} optimization task(s) failed")


def run_single_trip_batch(
    repo: Path,
    binary: Path,
    profile_root: Path,
    work_root: Path,
    tasks: list[dict[str, object]],
    output: Path,
) -> None:
    output.mkdir(parents=True, exist_ok=True)
    rows: list[dict[str, object]] = []
    failures = []
    for dataset in sorted({str(task["dataset"]) for task in tasks}):
        build_dataset_profiles(repo, profile_root, work_root, dataset)
    for position, task in enumerate(tasks, start=1):
        task_id = str(task["task_id"])
        dataset = str(task["dataset"])
        profile = str(task["profile"])
        hour = int(task["start_hour"])
        seed = int(task["seed"])
        task_dir = output / task_id
        task_dir.mkdir(parents=True, exist_ok=True)
        print(f"[{position}/{len(tasks)}] optimizing {task_id}", flush=True)
        try:
            command = solver_command(
                binary,
                source_paths(repo, dataset),
                speed_path(work_root, dataset, profile),
                hour,
                seed,
            )
            command.append("--truck-single-trip")
            _, elapsed, _ = run_command(
                command,
                task_dir,
                task_dir / "solver.log",
                TIME_LIMIT_SECONDS + 120,
                (0,),
            )
            solution = task_dir / "output_solution_best.txt"
            saved_solution = task_dir / f"final_solution_{task_id}.txt"
            shutil.copy2(solution, saved_solution)
            parsed = parse_solution(saved_solution)
            metrics = parse_route_metrics(str(parsed["final_solution_text"]))
            rows.append(
                {
                    "task_id": task_id,
                    "variant": "ST",
                    "dataset": dataset,
                    "customers": 100,
                    "actual_profile": profile,
                    "start_hour": hour,
                    "seed": seed,
                    "feasibility": parsed["feasibility"],
                    "makespan_s": parsed["makespan_s"],
                    "makespan_min": (
                        float(parsed["makespan_s"]) / 60.0
                        if parsed["makespan_s"] is not None
                        else None
                    ),
                    **metrics,
                    "total_truck_waiting_time_s": None,
                    "cpu_wall_time_s": elapsed,
                    "max_iterations": MAX_ITERATIONS,
                    "time_limit_s": TIME_LIMIT_SECONDS,
                    "solution_file": str(saved_solution.relative_to(output)),
                }
            )
            write_csv(output / "single_trip_results.csv", rows)
        except Exception as error:
            failures.append({"task_id": task_id, "error": repr(error)})
            (task_dir / "failure.txt").write_text(repr(error) + "\n", encoding="utf-8")
            print(f"FAILED {task_id}: {error}", flush=True)
    (output / "failures.json").write_text(
        json.dumps(failures, indent=2) + "\n", encoding="utf-8"
    )
    if failures:
        raise SystemExit(f"{len(failures)} single-trip task(s) failed")


def add_sheet(workbook, name: str, rows: list[dict[str, object]]) -> None:
    from openpyxl.styles import Font, PatternFill

    sheet = workbook.create_sheet(name)
    if not rows:
        sheet.append(["No data"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def read_unique_rows(input_root: Path, filename: str, key: str) -> list[dict[str, str]]:
    unique = {}
    for path in input_root.rglob(filename):
        with path.open(encoding="utf-8", newline="") as stream:
            for row in csv.DictReader(stream):
                unique[row[key]] = row
    return list(unique.values())


def aggregate(input_root: Path, output: Path, policy: str | None = None) -> None:
    from openpyxl import Workbook

    optimization = read_unique_rows(input_root, "optimization_results.csv", "task_id")
    evaluations = read_unique_rows(input_root, "evaluation_results.csv", "evaluation_id")
    optimization.sort(key=lambda row: row["task_id"])
    evaluations.sort(key=lambda row: row["evaluation_id"])
    if policy and any(row["policy"] != policy for row in [*optimization, *evaluations]):
        raise SystemExit(f"Unexpected policy in {policy} results")

    p2_by_case = {
        (row["dataset"], row["actual_profile"], row["start_hour"], row["seed"]): row
        for row in evaluations
        if row["policy"] == "P2"
    }
    summary_groups: dict[tuple[str, str, str], list[dict[str, str]]] = defaultdict(list)
    regrets: dict[tuple[str, str, str], list[float]] = defaultdict(list)
    for row in evaluations:
        group = (row["policy"], row["actual_profile"], row["start_hour"])
        summary_groups[group].append(row)
        if row["policy"] in ("P0", "P1") and row["feasibility"] == "FEASIBLE":
            benchmark = p2_by_case.get(
                (row["dataset"], row["actual_profile"], row["start_hour"], row["seed"])
            )
            if benchmark and benchmark["feasibility"] == "FEASIBLE":
                actual = float(row["realized_makespan_s"])
                perfect = float(benchmark["realized_makespan_s"])
                regrets[group].append(100.0 * (actual - perfect) / perfect)

    summary_rows = []
    for group, rows in sorted(summary_groups.items()):
        feasible = sum(row["feasibility"] == "FEASIBLE" for row in rows)
        makespans = [
            float(row["realized_makespan_s"])
            for row in rows
            if row["feasibility"] == "FEASIBLE" and row["realized_makespan_s"]
        ]
        summary_rows.append(
            {
                "policy": group[0],
                "actual_profile": group[1],
                "start_hour": group[2],
                "runs": len(rows),
                "feasible": feasible,
                "failure_rate_pct": 100.0 * (len(rows) - feasible) / len(rows),
                "mean_feasible_makespan_s": statistics.fmean(makespans) if makespans else None,
                "conditional_regret_pct": statistics.fmean(regrets[group]) if regrets[group] else None,
                "regret_cases": len(regrets[group]),
            }
        )

    solution_rows = []
    for path in sorted(input_root.rglob("final_solution_*.txt")):
        solution_rows.append(
            {
                "solution_file": path.name,
                "artifact_path": str(path.relative_to(input_root)),
                "final_solution_text": path.read_text(encoding="utf-8", errors="replace")[:32767],
            }
        )
    expected_tasks = policy_tasks()[policy] if policy else [task for tasks in policy_tasks().values() for task in tasks]
    expected_evaluations = INSTANCE_COUNT * 120 if policy else INSTANCE_COUNT * 360
    expected_evaluation_ids = {
        f"{p}-{dataset}-{profile}-{hour:02d}h-seed{seed:02d}"
        for p in ((policy,) if policy else ("P0", "P1", "P2"))
        for dataset in DATASETS for profile in PROFILES for hour in START_HOURS for seed in SEEDS
    }
    validation = [
        {"item": "optimization_runs", "found": len(optimization), "expected": len(expected_tasks)},
        {"item": "realized_evaluations", "found": len(evaluations), "expected": expected_evaluations},
        {"item": "saved_final_solutions", "found": len(solution_rows), "expected": len(expected_tasks)},
        {"item": "expected_task_ids", "found": len({row['task_id'] for row in optimization} & {str(task['task_id']) for task in expected_tasks}), "expected": len(expected_tasks)},
        {"item": "expected_evaluation_ids", "found": len({row['evaluation_id'] for row in evaluations} & expected_evaluation_ids), "expected": expected_evaluations},
    ]
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "Summary", summary_rows)
    add_sheet(workbook, "Optimization", optimization)
    add_sheet(workbook, "Evaluations", evaluations)
    add_sheet(workbook, "Final_Solutions", solution_rows)
    add_sheet(workbook, "Validation", validation)
    workbook.save(output)
    complete = all(row["found"] == row["expected"] for row in validation)
    print(json.dumps({"workbook": str(output), "complete": complete, "validation": validation}))
    if not complete:
        raise SystemExit("Aggregate is incomplete")


def aggregate_single_trip(input_root: Path, output: Path) -> None:
    from openpyxl import Workbook

    rows = read_unique_rows(input_root, "single_trip_results.csv", "task_id")
    rows.sort(
        key=lambda row: (
            row["dataset"], row["actual_profile"], int(row["start_hour"]), int(row["seed"])
        )
    )
    groups: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        groups[(row["actual_profile"], row["start_hour"])].append(row)
    summary_rows = []
    for (profile, hour), group in sorted(groups.items()):
        feasible_rows = [row for row in group if row["feasibility"] == "FEASIBLE"]
        makespans = [float(row["makespan_min"]) for row in feasible_rows if row["makespan_min"]]
        summary_rows.append(
            {
                "customers": 100,
                "day_profile": profile,
                "start_hour": hour,
                "runs": len(group),
                "feasible_runs": len(feasible_rows),
                "feasible_solution_rate_pct": 100.0 * len(feasible_rows) / len(group),
                "mean_feasible_makespan_min": statistics.fmean(makespans) if makespans else None,
                "mean_cpu_wall_time_s": statistics.fmean(float(row["cpu_wall_time_s"]) for row in group),
                "mean_drone_customers_pct": statistics.fmean(float(row["drone_customers_pct"]) for row in feasible_rows) if feasible_rows else None,
            }
        )
    solution_rows = []
    for path in sorted(input_root.rglob("final_solution_*.txt")):
        solution_rows.append(
            {
                "solution_file": path.name,
                "artifact_path": str(path.relative_to(input_root)),
                "final_solution_text": path.read_text(encoding="utf-8", errors="replace")[:32767],
            }
        )
    validation = [
        {"item": "single_trip_runs", "found": len(rows), "expected": INSTANCE_COUNT * len(PROFILES) * len(START_HOURS) * len(SEEDS)},
        {"item": "saved_final_solutions", "found": len(solution_rows), "expected": INSTANCE_COUNT * len(PROFILES) * len(START_HOURS) * len(SEEDS)},
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "Summary", summary_rows)
    add_sheet(workbook, "Single_Trip_Runs", rows)
    add_sheet(workbook, "Final_Solutions", solution_rows)
    add_sheet(workbook, "Validation", validation)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    complete = all(row["found"] == row["expected"] for row in validation)
    print(json.dumps({"workbook": str(output), "complete": complete, "validation": validation}))
    if not complete:
        raise SystemExit("Single-trip aggregate is incomplete")


def main() -> None:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)
    matrix_parser = subparsers.add_parser("matrix")
    matrix_parser.add_argument("--policy", choices=("P0", "P1", "P2"))
    subparsers.add_parser("single-trip-matrix")
    validate = subparsers.add_parser("validate-data")
    validate.add_argument("--repo", type=Path, default=Path.cwd())
    validate.add_argument("--profile-root", type=Path, required=True)
    build = subparsers.add_parser("build-dataset-profiles")
    build.add_argument("--repo", type=Path, default=Path.cwd())
    build.add_argument("--profile-root", type=Path, required=True)
    build.add_argument("--work-root", type=Path, required=True)
    build.add_argument("--dataset", required=True, choices=DATASETS)
    build.add_argument("--aggregation", choices=("arithmetic", "harmonic"), default="arithmetic")
    run = subparsers.add_parser("run-batch")
    run.add_argument("--repo", type=Path, default=Path.cwd())
    run.add_argument("--binary", type=Path, required=True)
    run.add_argument("--profile-root", type=Path, required=True)
    run.add_argument("--work-root", type=Path, required=True)
    run.add_argument("--tasks-json", required=True)
    run.add_argument("--output", type=Path, required=True)
    run.add_argument("--aggregation", choices=("arithmetic", "harmonic"), default="arithmetic")
    run.add_argument("--iterations", type=int, default=MAX_ITERATIONS)
    run.add_argument("--segment-iters", type=int, default=SEGMENT_ITERATIONS)
    run.add_argument("--time-limit", type=int, default=TIME_LIMIT_SECONDS)
    run_st = subparsers.add_parser("run-single-trip-batch")
    run_st.add_argument("--repo", type=Path, default=Path.cwd())
    run_st.add_argument("--binary", type=Path, required=True)
    run_st.add_argument("--profile-root", type=Path, required=True)
    run_st.add_argument("--work-root", type=Path, required=True)
    run_st.add_argument("--tasks-json", required=True)
    run_st.add_argument("--output", type=Path, required=True)
    aggregate_parser = subparsers.add_parser("aggregate")
    aggregate_parser.add_argument("--input-root", type=Path, required=True)
    aggregate_parser.add_argument("--output", type=Path, required=True)
    aggregate_parser.add_argument("--policy", choices=("P0", "P1", "P2"))
    aggregate_st = subparsers.add_parser("aggregate-single-trip")
    aggregate_st.add_argument("--input-root", type=Path, required=True)
    aggregate_st.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    if args.command == "matrix":
        print(json.dumps(build_matrix(args.policy), separators=(",", ":")))
    elif args.command == "single-trip-matrix":
        print(json.dumps(build_single_trip_matrix(), separators=(",", ":")))
    elif args.command == "validate-data":
        validate_data(args.repo.resolve(), args.profile_root.resolve())
    elif args.command == "build-dataset-profiles":
        build_dataset_profiles(
            args.repo.resolve(), args.profile_root.resolve(), args.work_root.resolve(), args.dataset, args.aggregation
        )
    elif args.command == "run-batch":
        run_batch(
            args.repo.resolve(),
            args.binary.resolve(),
            args.profile_root.resolve(),
            args.work_root.resolve(),
            [expand_task(task) for task in json.loads(args.tasks_json)],
            args.output.resolve(),
            args.aggregation,
            args.iterations,
            args.segment_iters,
            args.time_limit,
        )
    elif args.command == "run-single-trip-batch":
        run_single_trip_batch(
            args.repo.resolve(),
            args.binary.resolve(),
            args.profile_root.resolve(),
            args.work_root.resolve(),
            [expand_single_trip_task(task) for task in json.loads(args.tasks_json)],
            args.output.resolve(),
        )
    elif args.command == "aggregate-single-trip":
        aggregate_single_trip(args.input_root.resolve(), args.output.resolve())
    else:
        aggregate(args.input_root.resolve(), args.output.resolve(), args.policy)


if __name__ == "__main__":
    main()
