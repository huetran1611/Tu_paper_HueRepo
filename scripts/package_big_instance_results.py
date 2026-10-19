"""Package the seven already-downloaded workflows without rerunning experiments."""
from collections import Counter, defaultdict
import csv
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
import json
import math
from pathlib import Path
import shutil
import statistics
import subprocess
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from audit_seven_truck_loads import demands, trips
from collect_seven_experiment_runs import OUTPUT as SOURCE, ROOT, RUN_IDS
from summarize_seven_experiment_runs import parse_solution

DESTINATION = ROOT / "result_final/Big_instance"
DIMENSIONS = ("experiment", "customers", "scenario", "mode", "drone_capacity_kg", "trucks", "drones", "variant")


def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def dump(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2) + "\n")


def summarize(rows):
    valid = [r for r in rows if r["eligible"] == "True"]
    values = [float(r["makespan_s"]) for r in valid]
    return {"selected_slots": len(rows), "distinct_source_runs": len({r["global_id"] for r in rows}),
        "replaced_slots": sum(r["replacement_applied"] == "True" for r in rows),
        "feasible_audited_slots": len(valid), "instances": len({r["instance"] for r in rows}),
        "best_makespan_s": min(values) if values else None,
        "mean_makespan_s": statistics.mean(values) if values else None,
        "worst_makespan_s": max(values) if values else None,
        "std_makespan_s": statistics.stdev(values) if len(values) > 1 else None,
        "mean_solver_elapsed_s": statistics.mean(float(r["solver_elapsed_s"]) for r in rows),
        "mean_drone_customer_pct": statistics.mean(float(r["drone_customer_pct"]) for r in valid),
        "mean_truck_trips": statistics.mean(int(r["truck_trips"]) for r in valid),
        "max_truck_trip_load_kg": str(max(Decimal(r["max_truck_trip_load_kg"]) for r in rows)),
        "source_truck_capacity_kg": ";".join(sorted({r["truck_capacity_kg"] for r in rows}, key=float))}


def grouped(rows, fields):
    groups = defaultdict(list)
    for row in rows:
        groups[tuple(row[field] for field in fields)].append(row)
    result = []
    for key, group in sorted(groups.items()):
        result.append({**dict(zip(fields, key)), **summarize(group)})
    return result


def compare(left, right, label, customers):
    a = defaultdict(list)
    b = defaultdict(list)
    for row in left:
        if row["eligible"] == "True":
            a[row["instance"]].append(float(row["makespan_s"]))
    for row in right:
        if row["eligible"] == "True":
            b[row["instance"]].append(float(row["makespan_s"]))
    if set(a) != set(b) or len(a) != 16:
        raise ValueError("Comparison lacks matching 16-instance coverage")
    ratios = [100 * (statistics.mean(a[k]) - statistics.mean(b[k])) / statistics.mean(a[k]) for k in sorted(a)]
    return {"comparison": label, "customers": customers, "paired_instances": len(a),
        "baseline_mean_makespan_s": statistics.mean(statistics.mean(v) for v in a.values()),
        "alternative_mean_makespan_s": statistics.mean(statistics.mean(v) for v in b.values()),
        "mean_instance_reduction_pct": statistics.mean(ratios),
        "alternative_wins": sum(r > 1e-8 for r in ratios),
        "ties": sum(abs(r) <= 1e-8 for r in ratios),
        "alternative_losses": sum(r < -1e-8 for r in ratios)}


def excel_value(field, value):
    if value == "":
        return None
    if value in ("True", "False"):
        return value == "True"
    if field == "source_truck_capacity_kg" and isinstance(value, str) and ";" in value:
        return value
    numeric = field.endswith(("_s", "_min", "_kg", "_pct", "_mps", "_J")) or field in {
        "customers", "run", "result_slot_run", "replacement_source_run", "solver_return_code",
        "trucks", "drones", "used_trucks", "truck_trips", "truck_trips_per_used_truck",
        "drone_customers", "grid_size", "attempts", "initial_solution_seed", "planned_iterations",
        "planned_segments", "iterations_per_segment", "no_improve_threshold", "actual_iterations",
        "max_load_truck", "max_load_trip", "max_load_customer_count", "source_sa_T0",
        "source_sa_alpha", "source_destroy_rate", "source_penalty_exponent", "source_penalty_increase",
        "source_penalty_decrease", "source_penalty_min", "source_penalty_max", "source_neighborhood_count",
        "source_gamma1", "source_gamma2", "source_gamma3", "source_gamma4", "drone_height_m",
        "drone_power_beta", "drone_power_gamma"}
    if numeric and isinstance(value, str):
        result = float(value)
        if not math.isfinite(result):
            raise ValueError("Invalid workbook numeric value")
        return int(result) if result.is_integer() else result
    return value


def workbook(path, tables, notes):
    book = Workbook()
    book.remove(book.active)
    for name, rows in [("Notes", [{"note": note} for note in notes]), *tables.items()]:
        sheet = book.create_sheet(name)
        fields = list(rows[0])
        sheet.append(fields)
        for row in rows:
            values = [excel_value(field, row.get(field)) for field in fields]
            if any(isinstance(value, str) and len(value) > 32767 for value in values):
                raise ValueError("Excel cell would truncate solution information")
            sheet.append(values)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="31667C")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, field in enumerate(fields, 1):
            column = sheet.cell(1, index).column_letter
            sheet.column_dimensions[column].width = 52 if "solution" in field or "routes" in field else 22
        if name == "Notes":
            sheet.column_dimensions["A"].width = 120
            for row in sheet.iter_rows(min_row=2):
                row[0].alignment = Alignment(wrap_text=True, vertical="top")
                sheet.row_dimensions[row[0].row].height = 42
    book.save(path)
    check = load_workbook(path, read_only=True)
    assert check["All_Solutions"].max_row == 1601
    assert check["Workflow_Coverage"].max_row == 8
    check.close()


def build(stage):
    selected = read_csv(SOURCE / "all_solutions_and_parameters.csv")
    original = read_csv(SOURCE / "all_solutions_and_parameters.original.csv")
    original_by_id = {r["global_id"]: r for r in original}
    assert len(selected) == len(original) == len(original_by_id) == 1600
    assert {r["result_slot_global_id"] for r in selected} == set(original_by_id)
    assert Counter(r["experiment"] for r in selected) == {"big": 320, "trip": 960, "vehicle": 320}
    source_files = {str(p.relative_to(SOURCE)): digest(p) for p in SOURCE.rglob("*") if p.is_file()}
    snapshot = stage / "provenance/source_snapshot"
    shutil.copytree(SOURCE, snapshot)
    downloads = json.loads((snapshot / "download_manifest.json").read_text())
    catalog = json.loads((snapshot / "catalog.json").read_text())
    tree = {r["path"]: r["sha"] for r in json.loads((snapshot / "source_tree.json").read_text())["tree"]
            if r["type"] == "blob"}
    assert len(downloads) == 1600 and all(r["status"] == "downloaded" for r in downloads)
    seen = set()
    for item in downloads:
        path = snapshot / item["file"]
        assert digest(path) == item["sha256"], str(path)
        with zipfile.ZipFile(path) as archive:
            task = json.loads(archive.read("task_manifest.json"))["task"]
            key = task["global_id"]
            assert key not in seen
            seen.add(key)
            row = original_by_id[key]
            assert int(row["github_run_id"]) == item["run_id"]
            names = [n for n in archive.namelist() if n.endswith("output_solution_best.txt")]
            assert len(names) == 1
            text = archive.read(names[0]).decode("utf-8")
            assert text == row["solution_text"], key
            parsed = parse_solution(text, int(row["customers"]))
            assert parsed["route_audit_ok"] and parsed["reported_feasibility"] == "FEASIBLE", key
            assert abs(parsed["makespan_s"] - float(row["makespan_s"])) < 1e-6
        if len(seen) % 400 == 0:
            print("Verified original artifacts", len(seen), "/1600", flush=True)
    assert seen == set(original_by_id)

    # Preserve exact historical base inputs, since current project paths have moved.
    weights = {}
    input_index = []
    for row in selected:
        name = row["input_file"]
        if name in weights:
            continue
        content = subprocess.check_output(["git", "show", f"{row['source_commit']}:{name}"], cwd=ROOT)
        blob = hashlib.sha1(b"blob " + str(len(content)).encode() + b"\0" + content).hexdigest()
        assert blob == tree[name], name
        path = stage / "data/inputs" / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        weights[name] = demands(path, int(row["customers"]))
        input_index.append({"original_project_path": name, "package_path": str(path.relative_to(stage)),
            "source_commit": row["source_commit"], "git_blob_sha": blob, "sha256": digest(path)})
    assert len(weights) == 96
    write_csv(stage / "data/input_index.csv", input_index)

    trip_loads = []
    (stage / "solutions").mkdir()
    for row in selected:
        source = original_by_id[row["global_id"]]
        assert all(row[k] == value for k, value in source.items()), row["global_id"]
        slot = row["result_slot_global_id"]
        local = stage / f"solutions/slot_{slot}.txt"
        local.write_text(row["solution_text"])
        per_case = []
        for truck, route in enumerate(json.loads(row["truck_routes"]), 1):
            for index, customers in enumerate(trips(route), 1):
                load = sum((weights[row["input_file"]][node] for node in customers), Decimal(0))
                per_case.append(load)
                trip_loads.append({"result_slot_global_id": slot, "source_global_id": row["global_id"],
                    "experiment": row["experiment"], "instance": row["instance"],
                    "truck": truck, "trip": index, "load_kg": str(load), "customers": len(customers),
                    "customer_ids": json.dumps(customers), "replacement_applied": row["replacement_applied"]})
        assert max(per_case, default=Decimal(0)) == Decimal(row["max_truck_trip_load_kg"]), slot
        assert all(load <= Decimal(40) for load in per_case), slot
        row.update(local_solution_file=str(local.relative_to(stage)),
                   local_artifact_file="provenance/source_snapshot/" + row["artifact_file"],
                   local_input_file="data/inputs/" + row["input_file"])
    write_csv(stage / "all_solutions_and_parameters.csv", selected)
    write_csv(stage / "report/truck_trip_loads.csv", trip_loads)

    families = {"capacity": [r for r in selected if r["experiment"] == "trip"],
                "single_multi_trip": [r for r in selected if r["experiment"] == "trip"],
                "vehicle_sweep": [r for r in selected if r["experiment"] == "vehicle"],
                "heavy_light": [r for r in selected if r["experiment"] == "big"]}
    configs = grouped(selected, DIMENSIONS)
    instances = grouped(selected, (*DIMENSIONS, "instance"))
    assert len(configs) == 20 and len(instances) == 320
    assert all(r["selected_slots"] == 5 for r in instances)
    write_csv(stage / "report/configuration_summary.csv", configs)
    write_csv(stage / "report/instance_configuration_summary.csv", instances)
    for name, rows in families.items():
        write_csv(stage / name / "solutions.csv", rows)
        write_csv(stage / name / "configuration_summary.csv", grouped(rows, DIMENSIONS))
    capacity, trip, fleet = [], [], []
    baseline = [r for r in families["capacity"] if r["mode"] == "multi_trip" and float(r["drone_capacity_kg"]) == 2.27]
    write_csv(stage / "vehicle_sweep/baseline_solutions.csv", baseline)
    for size in ("200", "500"):
        rows = [r for r in families["capacity"] if r["customers"] == size]
        for cap in (2.27, 5, 10):
            subset = [r for r in rows if float(r["drone_capacity_kg"]) == cap]
            trip.append(compare([r for r in subset if r["mode"] == "one_trip"],
                [r for r in subset if r["mode"] == "multi_trip"], f"single -> multi-trip; drone capacity {cap:g} kg", size))
        for mode in ("one_trip", "multi_trip"):
            subset = [r for r in rows if r["mode"] == mode]
            for cap in (5, 10):
                capacity.append(compare([r for r in subset if float(r["drone_capacity_kg"]) == 2.27],
                    [r for r in subset if float(r["drone_capacity_kg"]) == cap], f"{mode}; drone capacity 2.27 -> {cap} kg", size))
        variants = [r for r in families["vehicle_sweep"] if r["customers"] == size]
        for variant in sorted({r["variant"] for r in variants}):
            fleet.append(compare([r for r in baseline if r["customers"] == size],
                [r for r in variants if r["variant"] == variant], f"balanced fleet -> {variant}", size))
    for folder, rows in (("capacity", capacity), ("single_multi_trip", trip), ("vehicle_sweep", fleet)):
        write_csv(stage / folder / "comparisons.csv", rows)

    coverage = []
    for index, run in enumerate(catalog, 1):
        info = run["run"]
        rows = [r for r in selected if int(r["github_run_id"]) == info["id"]]
        expected = 250 if index < 7 else 100
        assert info["id"] == RUN_IDS[index - 1] and len(rows) == expected
        coverage.append({"batch": index, "github_run_id": info["id"], "workflow_name": info["name"],
            "workflow_path": info["path"], "source_commit": info["head_sha"], "expected_slots": expected,
            "selected_slots": len(rows), "distinct_selected_source_runs": len({r["global_id"] for r in rows}),
            "replaced_slots": sum(r["replacement_applied"] == "True" for r in rows),
            "source_url": info["html_url"]})
        path = stage / "provenance/workflows" / f"{info['id']}_{Path(info['path']).name}"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(subprocess.check_output(["git", "show", f"{info['head_sha']}:{info['path']}"], cwd=ROOT))
    write_csv(stage / "report/workflow_coverage.csv", coverage)
    notes = [
        "Offline package of seven downloaded workflows; no new optimization or replay.",
        "Selected table preserves the existing curated CSV: slot 0853 uses source run 0851. 1600 slots, 1599 distinct source runs.",
        "Source solver truck capacity was 1000 kg, NOT 40 kg. Verified selected trip loads are <=40 kg; this is not a rerun with capacity 40.",
        "capacity/ and single_multi_trip/ are different views of the SAME 960 trip/capacity records; do not add their counts.",
        "Capacity sweep changes DRONE capacity: 2.27, 5, 10 kg. Both single-trip and multi-trip modes are included.",
        "vehicle_sweep/ has 320 swept-fleet records plus 160 balanced-fleet baseline records reused from the trip experiment.",
        "200 customers: 5 trucks+5 drones baseline; 6+4 and 4+6 alternatives. 500: 10+10 baseline; 11+9 and 9+11 alternatives.",
        "heavy_light/ preserves the remaining 320 results from the same seven workflows.",
        "Repetitions 1-5 are slot labels, not distinct recorded seeds. Initialization seed 42; other RNGs time-based.",
        "Summaries use the selected 5 slots per instance/configuration, including the disclosed duplicate source. They are not 1600 independent runs.",
        "Makespan and solver elapsed times are seconds and are different quantities. Positive paired-instance reduction favors the alternative.",
        "All routes are embedded in CSV/Excel and also saved as solutions/slot_XXXX.txt; no GitHub access is needed to read them.",
        "Original artifact ZIPs and original audit tables are preserved under provenance/source_snapshot; those original audits are not the curated results.",
        "data/inputs has exact historical base instance text files; traffic matrices are referenced by original project path but not copied into this package."]
    workbook(stage / "report/Big_instance_results.xlsx", {"All_Solutions": selected,
        "Config_Summary": configs, "Instance_Config": instances, "Capacity_Comparison": capacity,
        "Trip_Comparison": trip, "Fleet_Comparison": fleet, "Workflow_Coverage": coverage}, notes)
    (stage / "README.md").write_text("# Big Instance Results\n\n" + "\n\n".join(notes) + "\n\n"
        "## Layout\n\n- all_solutions_and_parameters.csv: 1600 selected slots with full routes and parameters.\n"
        "- capacity/: 960 records, capacity summaries and paired comparisons.\n"
        "- single_multi_trip/: the same 960 records, single/multi-trip paired comparisons.\n"
        "- vehicle_sweep/: 320 records and 160 explicitly separated baseline records.\n"
        "- heavy_light/: 320 heavy/light demand scenario records.\n"
        "- solutions/: one readable solution file per selected slot.\n"
        "- report/: combined Excel, configuration/instance/workflow summaries, selected truck-trip loads.\n"
        "- data/: 96 historical base input text files and their Git blob/hash index.\n"
        "- provenance/: unmodified downloaded source snapshot and seven historical workflow YAML files.\n"
        "- manifest.json: content hashes and verification scope.\n")
    for name, rows in families.items():
        (stage / name / "README.md").write_text(f"# {name}\n\nSelected records: {len(rows)}.\n\n"
            "Use local_solution_file and local_artifact_file relative to the Big_instance package root. "
            "Full route text is also embedded in solutions.csv. See ../README.md for curation, capacity, and repetition caveats.\n")
    assert all(digest(SOURCE / path) == sha for path, sha in source_files.items())
    assert all(digest(snapshot / path) == sha for path, sha in source_files.items())
    packed = read_csv(stage / "all_solutions_and_parameters.csv")
    assert len(packed) == 1600
    for old, new in zip(read_csv(SOURCE / "all_solutions_and_parameters.csv"), packed):
        assert all(new[k] == value for k, value in old.items())
        assert (stage / new["local_solution_file"]).read_text() == new["solution_text"]
        assert (stage / new["local_artifact_file"]).is_file()
    return {"created_utc": datetime.now(timezone.utc).isoformat(), "source_directory": str(SOURCE),
        "selected_slots": 1600, "distinct_selected_source_runs": 1599, "original_artifacts": 1600,
        "workflow_ids": list(RUN_IDS), "view_counts": {k: len(v) for k, v in families.items()},
        "overlapping_views": ["capacity", "single_multi_trip"], "vehicle_baseline_records": len(baseline),
        "source_truck_capacity_kg": 1000, "selected_max_truck_load_kg": str(max(Decimal(r["max_truck_trip_load_kg"]) for r in selected)),
        "historical_input_files": len(weights), "configuration_groups": len(configs),
        "instance_configuration_groups": len(instances), "original_source_files_unchanged": True,
        "source_file_sha256": source_files, "new_optimizations": 0, "new_replays": 0,
        "files": [{"path": str(p.relative_to(stage)), "bytes": p.stat().st_size, "sha256": digest(p)}
                  for p in sorted(stage.rglob("*")) if p.is_file()]}


def main():
    if DESTINATION.exists():
        raise FileExistsError(f"Refusing to overwrite {DESTINATION}")
    stage = Path(tempfile.mkdtemp(prefix=".Big_instance-", dir=DESTINATION.parent))
    try:
        manifest = build(stage)
        dump(stage / "manifest.json", manifest)
        stage.rename(DESTINATION)
    except BaseException:
        shutil.rmtree(stage)
        raise
    print(json.dumps({k: v for k, v in manifest.items() if k not in ("source_file_sha256", "files")}, indent=2))
    print("Created", DESTINATION, flush=True)


if __name__ == "__main__":
    main()
