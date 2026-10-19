"""Audit and aggregate all 1,600 solver results from the seven requested runs."""
from collections import Counter, defaultdict
import argparse
import csv
import hashlib
import importlib.util
import json
import math
from pathlib import Path
import re
import statistics
import zipfile
from functools import lru_cache

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from collect_seven_experiment_runs import OUTPUT, REPO, RUN_IDS


SHA = "902415a453ae5174403151323de4398f7d456c6d"
SOURCE = OUTPUT / "sources" / SHA / "scripts/experiment_batches"
NUMBER = r"[-+0-9.eE]+"


def write_csv(name, rows):
    if not rows:
        return
    with (OUTPUT / name).open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def value(text, label):
    found = re.findall(r"^" + re.escape(label) + r":\s*(.+)$", text, re.MULTILINE | re.IGNORECASE)
    return found[-1].strip() if found else None


def number(text, label):
    found = value(text, label)
    return float(found.split()[0]) if found else None


def parse_routes(text):
    parsed = []
    for match in re.finditer(r"^(Truck|Drone) (\d+):\s*([\d\s]+)\|(?:Truck|Drone) Time:\s*(" + NUMBER + r")", text, re.MULTILINE):
        kind, vehicle, route, duration = match.groups()
        parsed.append((kind, int(vehicle), [int(v) for v in route.split()], float(duration)))
    return parsed


def trip_count(route):
    return sum(bool(a) for a in split_trips(route))


def split_trips(route):
    trips, current = [], []
    for node in route:
        if node:
            current.append(node)
        elif current:
            trips.append(current)
            current = []
    if current:
        trips.append(current)
    return trips


def parse_solution(text, customers):
    objective = number(text, "Improved solution cost")
    initial = number(text, "Initial solution cost")
    elapsed = number(text, "Mean elapsed time")
    if elapsed is None:
        elapsed = number(text, "Elapsed time")
    feasible = value(text, "Final solution feasibility")
    routes = parse_routes(text)
    issues = []
    visits = Counter(node for _, _, route, _ in routes for node in route if node)
    if visits != Counter(range(1, customers+1)):
        issues.append("customer_coverage_or_duplicate")
    if not routes or any(not route or route[0] != 0 or route[-1] != 0 for _, _, route, _ in routes):
        issues.append("route_endpoints")
    if objective is None or not math.isfinite(objective) or objective <= 0:
        issues.append("invalid_objective")
    elif routes and abs(max(duration for _, _, _, duration in routes)-objective) > 1e-4:
        issues.append("makespan_differs_from_max_vehicle_time")
    if elapsed is None or not math.isfinite(elapsed) or elapsed < 0:
        issues.append("invalid_elapsed")
    trucks = [r for r in routes if r[0] == "Truck"]
    drones = [r for r in routes if r[0] == "Drone"]
    drone_customers = sum(len([v for v in r[2] if v]) for r in drones)
    truck_trips = sum(trip_count(r[2]) for r in trucks)
    used_trucks = sum(any(r[2]) for r in trucks)
    return {"initial_makespan_s": initial, "makespan_s": objective,
            "makespan_min": objective / 60 if objective is not None else None,
            "solver_elapsed_s": elapsed, "reported_feasibility": feasible,
            "route_audit_ok": not issues, "route_audit_issues": ";".join(issues),
            "trucks": len(trucks), "drones": len(drones), "used_trucks": used_trucks,
            "truck_trips": truck_trips, "truck_trips_per_used_truck": truck_trips / used_trucks if used_trucks else 0,
            "drone_customers": drone_customers, "drone_customer_pct": 100 * drone_customers / customers,
            "truck_completion_s": max((r[3] for r in trucks), default=0),
            "drone_completion_s": max((r[3] for r in drones), default=0),
            "improvement_from_initial_pct": 100*(initial-objective)/initial if initial and objective else None,
            "route_fingerprint": hashlib.sha256(json.dumps(routes).encode()).hexdigest(),
            "truck_routes": json.dumps([r[2] for r in trucks]), "drone_routes": json.dumps([r[2] for r in drones])}


@lru_cache(maxsize=2)
def source_settings(mode):
    filename = "tabubu_TimeDependent_1trip.cpp" if mode == "one_trip" else "tabubu_TimeDependent.cpp"
    source = (OUTPUT / "sources" / SHA / filename).read_text()
    uncommented = re.sub(r"/\*.*?\*/|//[^\n]*", "", source, flags=re.DOTALL)
    names = {"Dh": "truck_capacity_kg", "E": "drone_energy_capacity_J",
             "v_fly_drone": "drone_speed_mps", "v_take_off": "drone_takeoff_speed_mps",
             "v_landing": "drone_landing_speed_mps", "height": "drone_height_m",
             "power_beta": "drone_power_beta", "power_gamma": "drone_power_gamma",
             "gamma1": "source_gamma1", "gamma2": "source_gamma2", "gamma3": "source_gamma3",
             "gamma4": "source_gamma4", "T0": "source_sa_T0", "alpha": "source_sa_alpha",
             "DESTROY_RATE": "source_destroy_rate", "PENALTY_EXPONENT": "source_penalty_exponent",
             "PENALTY_INCREASE": "source_penalty_increase", "PENALTY_DECREASE": "source_penalty_decrease",
             "PENALTY_MIN": "source_penalty_min", "PENALTY_MAX": "source_penalty_max",
             "NUM_NEIGHBORHOODS": "source_neighborhood_count"}
    result = {}
    for variable, column in names.items():
        found = re.search(r"\b"+re.escape(variable)+r"\s*=\s*("+NUMBER+r")", uncommented)
        result[column] = float(found[1]) if found else None
    result["solver_source_file"] = filename
    result["solver_source_sha256"] = hashlib.sha256(source.encode()).hexdigest()
    return result


def mean(values):
    values = [v for v in values if v is not None]
    return statistics.fmean(values) if values else None


def group_key(row):
    return (row["experiment"], row["customers"], row["scenario"], row["mode"],
            row["drone_capacity_kg"], row["trucks"], row["drones"], row["variant"])


def group_labels(key):
    return dict(zip(("experiment", "customers", "scenario", "mode", "drone_capacity_kg",
                     "trucks", "drones", "variant"), key))


def aggregate_group(rows):
    valid = [r for r in rows if r["eligible"]]
    values = [r["makespan_s"] for r in valid]
    return {"runs": len(rows), "feasible_audited_runs": len(valid),
            "feasible_rate_pct": 100*len(valid)/len(rows),
            "best_makespan_s": min(values) if values else None,
            "mean_makespan_s": mean(values), "worst_makespan_s": max(values) if values else None,
            "std_makespan_s": statistics.stdev(values) if len(values)>1 else None,
            "mean_makespan_min": mean(values)/60 if values else None,
            "mean_solver_elapsed_s": mean([r["solver_elapsed_s"] for r in rows]),
            "mean_drone_customer_pct": mean([r["drone_customer_pct"] for r in valid]),
            "mean_truck_trips": mean([r["truck_trips"] for r in valid]),
            "mean_trips_per_used_truck": mean([r["truck_trips_per_used_truck"] for r in valid]),
            "mean_initial_improvement_pct": mean([r["improvement_from_initial_pct"] for r in valid])}


def compare_instances(left, right, label, customers, capacity=None):
    a = {r["instance"]: r for r in left if r["mean_makespan_s"] is not None}
    b = {r["instance"]: r for r in right if r["mean_makespan_s"] is not None}
    common = sorted(set(a) & set(b))
    reductions = [100*(a[k]["mean_makespan_s"]-b[k]["mean_makespan_s"])/a[k]["mean_makespan_s"] for k in common]
    differences = [a[k]["mean_makespan_s"]-b[k]["mean_makespan_s"] for k in common]
    return {"comparison": label, "customers": customers, "drone_capacity_kg": capacity,
            "paired_instances": len(common), "baseline_mean_s": mean([a[k]["mean_makespan_s"] for k in common]),
            "alternative_mean_s": mean([b[k]["mean_makespan_s"] for k in common]),
            "mean_instance_reduction_pct": mean(reductions),
            "alternative_wins": sum(d>1e-6 for d in differences), "ties": sum(abs(d)<=1e-6 for d in differences),
            "alternative_losses": sum(d< -1e-6 for d in differences)}


def workbook(tables, notes):
    book = Workbook()
    book.remove(book.active)
    for name, rows in tables.items():
        sheet = book.create_sheet(name)
        if not rows:
            continue
        sheet.append(list(rows[0]))
        for row in rows:
            sheet.append([json.dumps(v) if isinstance(v, (dict, list)) else v for v in row.values()])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="36534B")
        for column in sheet.columns:
            label = str(column[0].value)
            sheet.column_dimensions[column[0].column_letter].width = min(42, max(15, len(label)+2))
            for cell in column[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"
                elif isinstance(cell.value, str) and cell.value.startswith("https://github.com/"):
                    cell.hyperlink = cell.value
    sheet = book.create_sheet("Notes")
    for note in notes:
        sheet.append([note])
    sheet.column_dimensions["A"].width = 110
    book.save(OUTPUT / "seven_experiments_summary.xlsx")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--csv-only", action="store_true")
    args = parser.parse_args()
    spec = importlib.util.spec_from_file_location("original_experiment_plan", SOURCE / "plan.py")
    plan = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(plan)
    expected = {t["global_id"]: t for t in plan.all_tasks()}
    digest = plan.plan_hash(list(expected.values()))
    downloads = json.loads((OUTPUT / "download_manifest.json").read_text())
    if len(downloads) != 1600 or any(d["status"] != "downloaded" for d in downloads):
        raise ValueError("All 1600 artifact downloads must complete first")
    all_rows, manifests, issues = [], [], []
    seen = set()
    for item in downloads:
        path = OUTPUT / item["file"]
        if hashlib.sha256(path.read_bytes()).hexdigest() != item["sha256"]:
            raise ValueError(f"Artifact hash changed: {path}")
        with zipfile.ZipFile(path) as archive:
            manifest = json.loads(archive.read("task_manifest.json"))
            task = manifest["task"]
            key = task["global_id"]
            if key in seen or task != expected.get(key) or manifest["plan_hash"] != digest:
                raise ValueError(f"Duplicate or unexpected task {key}")
            batch = (int(key)-1)//250+1
            if manifest["batch"] != batch or item["run_id"] != RUN_IDS[batch-1] or manifest["commit"] != SHA:
                raise ValueError(f"Provenance mismatch for {key}")
            seen.add(key)
            manifests.append(manifest)
            solutions = [name for name in archive.namelist() if name.endswith("output_solution_best.txt")]
            if len(solutions) != 1:
                raise ValueError(f"Expected exactly one primary solution in {path}")
            text = archive.read(solutions[0]).decode("utf-8", errors="replace")
            customers = int(task.get("number_of_customer", task.get("customer_count")))
            parsed = parse_solution(text, customers)
            log_names = [n for n in archive.namelist() if n.endswith("run.log") or n.endswith("solver_stdout_stderr.txt")]
            log = archive.read(log_names[0]).decode("utf-8", errors="replace") if log_names else ""
            config = re.search(r"Search config: total_iters=(\d+) \(segments=(\d+), iters_per_seg=(\d+), no_improve=(\d+)\)", log)
            iteration_names = [n for n in archive.namelist() if n.endswith("output_mode_0.txt")]
            actual_iterations = None
            if iteration_names:
                lines = archive.read(iteration_names[0]).decode().splitlines()
                if len(lines) > 1:
                    actual_iterations = int(next(csv.DictReader([lines[0], lines[-1]]))["iter"])
            elapsed_process = None
            if task["experiment"] == "trip":
                result_name = next(n for n in archive.namelist() if n.endswith("/result.json"))
                source_result = json.loads(archive.read(result_name))
                elapsed_process = float(source_result["elapsed_seconds"])
                if abs(float(source_result["best_objective"])-parsed["makespan_s"]) > 1e-5:
                    parsed["route_audit_ok"] = False
                    parsed["route_audit_issues"] += ";result_json_objective_mismatch"
            row = {"global_id": key, "github_run_id": item["run_id"], "batch": batch,
                   "experiment": task["experiment"], "instance": task["base"], "customers": customers,
                   "scenario": task.get("scenario", ""), "mode": task.get("mode", "multi_trip"),
                   "drone_capacity_kg": float(task.get("drone_capacity", 2.27)),
                   "variant": task.get("variant", ""), "run": task["run"],
                   "solver_return_code": manifest["return_code"], "time_limit_s": float(manifest["time_limit_sec"]),
                   **parsed, "process_elapsed_s": elapsed_process,
                   "eligible": manifest["return_code"] == 0 and parsed["reported_feasibility"] == "FEASIBLE" and parsed["route_audit_ok"],
                   "source_url": f"https://github.com/{REPO}/actions/runs/{item['run_id']}",
                   "artifact_id": item["artifact_id"], "artifact_file": item["file"],
                   "source_commit": SHA, "plan_hash": digest,
                   "input_file": task["file"], "truck_vmax_file": task["vmax_file"], "truck_theta_file": task["theta_file"],
                   "input_weight_label": task.get("weight"), "grid_size": task.get("gridsize", int(task["base"].split(".")[1])),
                   "attempts": 1, "explicit_run_seed": None, "initial_solution_seed": 42,
                   "rng_note": "No per-run --seed; initialization seed 42, other RNGs time-based",
                   "planned_iterations": int(config[1]) if config else None,
                   "planned_segments": int(config[2]) if config else None,
                   "iterations_per_segment": int(config[3]) if config else None,
                   "no_improve_threshold": int(config[4]) if config else None,
                   "actual_iterations": actual_iterations,
                   "completed_iteration_budget": actual_iterations == int(config[1]) if config and actual_iterations is not None else None,
                   "command_as_logged": value(log, "Command"),
                   "task_manifest_json": json.dumps(manifest, separators=(",", ":")),
                   "solution_text": text,
                   **source_settings(task.get("mode", "multi_trip"))}
            all_rows.append(row)
            if not row["eligible"]:
                issues.append({k: row[k] for k in ("global_id", "experiment", "instance", "reported_feasibility", "route_audit_issues", "solver_return_code")})
    if seen != set(expected) or len({m["time_limit_sec"] for m in manifests}) != 1:
        raise ValueError("Incomplete coverage or mixed search time limits")
    all_rows.sort(key=lambda row: row["global_id"])
    if args.csv_only:
        if any(r["planned_iterations"] is None or r["actual_iterations"] is None for r in all_rows):
            raise ValueError("Missing search configuration or iteration history")
        families = Counter(r["experiment"] for r in all_rows)
        if families != {"big": 320, "trip": 960, "vehicle": 320}:
            raise ValueError("Incorrect experiment-family coverage")
        repetitions = defaultdict(set)
        for row in all_rows:
            key = (group_key(row), row["instance"])
            repetitions[key].add(row["run"])
        if len(repetitions) != 320 or any(r != {1, 2, 3, 4, 5} for r in repetitions.values()):
            raise ValueError("Incorrect five-repetition coverage")
        filename = "all_solutions_and_parameters.csv"
        write_csv(filename, all_rows)
        with (OUTPUT / filename).open(newline="", encoding="utf-8-sig") as handle:
            roundtrip = list(csv.DictReader(handle))
        assert len(roundtrip) == 1600 and len({r["global_id"] for r in roundtrip}) == 1600
        assert all(r["solution_text"] and r["truck_routes"] and r["drone_routes"] for r in roundtrip)
        print(json.dumps({"file": str(OUTPUT / filename), "rows": len(all_rows), "columns": len(all_rows[0]),
                          "experiment_counts": families, "reported_feasible": sum(r["reported_feasibility"] == "FEASIBLE" for r in all_rows),
                          "feasible_and_route_audited": sum(r["eligible"] for r in all_rows),
                          "audit_issues": issues, "five_runs_per_configuration_verified": True,
                          "csv_roundtrip_verified": True}, indent=2))
        return
    groups, instance_groups = defaultdict(list), defaultdict(list)
    for row in all_rows:
        groups[group_key(row)].append(row)
        instance_groups[(group_key(row), row["instance"])].append(row)
    configs, instances = [], []
    for (key, instance), rows in sorted(instance_groups.items()):
        if len(rows) != 5 or {r["run"] for r in rows} != set(range(1, 6)):
            raise ValueError(f"Expected five distinct repetitions for {key}/{instance}")
        valid = [r for r in rows if r["eligible"]]
        best = min(valid, key=lambda r: r["makespan_s"], default=None)
        instances.append({**group_labels(key), "instance": instance, **aggregate_group(rows),
                          "best_run": best["run"] if best else None,
                          "distinct_route_solutions": len({r["route_fingerprint"] for r in valid})})
    for key, rows in sorted(groups.items()):
        corresponding = [r for r in instances if group_key(r) == key]
        configs.append({**group_labels(key), "instance_count": len(corresponding), **aggregate_group(rows),
                        "mean_best_of_5_s": mean([r["best_makespan_s"] for r in corresponding])})
    trip_comparisons, capacity_comparisons, fleet_comparisons = [], [], []
    for customers in (200, 500):
        subset = [r for r in instances if r["customers"] == customers]
        for capacity in (2.27, 5., 10.):
            base = [r for r in subset if r["experiment"]=="trip" and r["mode"]=="one_trip" and r["drone_capacity_kg"]==capacity]
            alternative = [r for r in subset if r["experiment"]=="trip" and r["mode"]=="multi_trip" and r["drone_capacity_kg"]==capacity]
            trip_comparisons.append(compare_instances(base, alternative, "one_trip -> multi_trip", customers, capacity))
        for mode in ("multi_trip", "one_trip"):
            base = [r for r in subset if r["experiment"]=="trip" and r["mode"]==mode and r["drone_capacity_kg"]==2.27]
            for capacity in (5., 10.):
                alt = [r for r in subset if r["experiment"]=="trip" and r["mode"]==mode and r["drone_capacity_kg"]==capacity]
                capacity_comparisons.append(compare_instances(base, alt, f"{mode}: Dd 2.27 -> {capacity}", customers, capacity))
        base = [r for r in subset if r["experiment"]=="trip" and r["mode"]=="multi_trip" and r["drone_capacity_kg"]==2.27]
        for variant in sorted({r["variant"] for r in subset if r["experiment"]=="vehicle"}):
            alt = [r for r in subset if r["experiment"]=="vehicle" and r["variant"]==variant]
            fleet_comparisons.append(compare_instances(base, alt, f"balanced baseline -> {variant}", customers, 2.27))
    coverage = []
    for index, run_id in enumerate(RUN_IDS, 1):
        rows = [r for r in all_rows if r["github_run_id"] == run_id]
        coverage.append({"batch": index, "github_run_id": run_id, "expected": 250 if index<7 else 100,
                         "actual": len(rows), "feasible_audited": sum(r["eligible"] for r in rows),
                         "source_url": f"https://github.com/{REPO}/actions/runs/{run_id}"})
    validation = {"expected_tasks": 1600, "actual_tasks": len(all_rows), "unique_tasks": len(seen),
                  "plan_hash": digest, "commit": SHA, "download_hashes_verified": True,
                  "five_repetitions_per_instance_configuration": True,
                  "configuration_groups": len(configs), "instance_configuration_groups": len(instances),
                  "reported_feasible": sum(r["reported_feasibility"]=="FEASIBLE" for r in all_rows),
                  "feasible_and_route_audited": sum(r["eligible"] for r in all_rows),
                  "anomalies": issues, "time_limit_s": all_rows[0]["time_limit_s"]}
    (OUTPUT / "validation.json").write_text(json.dumps(validation, indent=2)+"\n")
    tables = {"All_Runs": all_rows, "Config_Summary": configs, "Instance_Config": instances,
              "Trip_Comparison": trip_comparisons, "Capacity_Comparison": capacity_comparisons,
              "Fleet_Comparison": fleet_comparisons, "Run_Coverage": coverage, "Anomalies": issues}
    for name, rows in tables.items():
        write_csv(name.lower()+".csv", rows)
    notes = ["Objective/makespan and elapsed times are in seconds; makespan_min divides by 60.",
             "Makespan is a modeled completion duration, not solver running time.",
             "Objective summaries include only reported FEASIBLE results passing route coverage and makespan checks.",
             "mean_best_of_5_s averages each instance's best feasible solution, not the best instance in the group.",
             "Each configuration has 16 instances and five repetitions per instance.",
             "Comparisons pair instance means over repetitions; positive reduction favors the alternative.",
             "Repetition labels 1-5 are NOT recorded distinct seeds; runner commands do not pass --seed.",
             "Source solver uses initialization seed 42 and time-based random generators elsewhere.",
             "This audit checks saved routes and reported feasibility, not independent replay of physical constraints.",
             "Fleet baseline is trip experiment multi_trip at Dd=2.27 on matching instances; 200:5+5, 500:10+10.",
             "Source commit: "+SHA, "Plan hash: "+digest]
    workbook(tables, notes)
    print(json.dumps({"validation": validation, "trip_comparisons": trip_comparisons,
                      "fleet_comparisons": fleet_comparisons}, indent=2))


if __name__ == "__main__":
    main()
